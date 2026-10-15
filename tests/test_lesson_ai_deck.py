"""«Жоспар»: колода из шаблонов — каталог, исполнитель сценария, правка.

Без базы: каталог библиотеки, сборка `.pptx` из шаблонов-фикстур и обе
формы слайда презентации — чистые функции над содержимым материала.
Библиотека-фикстура собирается кодом (`tests/fixtures/pptx_templates`):
настоящую библиотеку делает дизайн параллельно, и исполнитель не должен
её ждать.
"""

import logging
import re
from io import BytesIO
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from app.adapters import lesson_files, lesson_preview
from app.adapters.docx.common import DocHeader
from app.adapters.pptx import deck
from app.lesson_ai import edit_tools, images, pptx_templates
from app.lesson_ai.blocks import BlockRewriter
from app.lesson_ai.kit_map import kit_outline
from app.lesson_ai.prompts import edit_chat
from app.lesson_ai.prompts import presentation as prompt
from app.lesson_ai.schemas import strict_schema
from app.lesson_ai.schemas.presentation_deck import (
    PresentationDeckDraft,
    any_form_draft,
)
from tests.fixtures.pptx_templates.build import build as build_library

HEADER = DocHeader(lang="ru", teacher_name="Нурланова Айгуль", grade=3, topic="Глагол to be")
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


@pytest.fixture(scope="module")
def lib(tmp_path_factory) -> Path:
    return build_library(tmp_path_factory.mktemp("deck_lib"))


def scenario_slide(template: str, **over) -> dict:
    return {"template": template, "fill": [], "remove": [], "add": [], "notes": "Заметка", **over}


def fill(slot: str, *, text: str = "", items: list | None = None, image: str = "") -> dict:
    return {"slot": slot, "text": text, "items": items or [], "image": image}


def question_slide(**over) -> dict:
    return scenario_slide(
        "question_choice",
        fill=[
            fill("question", text="Where is my bag?"),
            fill("options", items=["It is here", "It is under the chair"]),
            fill("answer", text="It is under the chair"),
            fill("picture", image="backpack"),
        ],
        **over,
    )


def deck_content(*slides: dict) -> dict:
    return {
        "goals": ["Научимся"],
        "slides_beginning": list(slides[:1]),
        "slides_middle": list(slides[1:]),
        "slides_end": [],
    }


def built(lib, *slides: dict):
    data = deck.render_deck(deck_content(*slides), HEADER, lib).data
    return data, Presentation(BytesIO(data))


def texts(sld) -> list[str]:
    return [
        p.text for s in sld.shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.text
    ]


def shapes_by_name(sld) -> dict:
    return {s.name: s for s in sld.shapes}


def pictures(sld) -> list:
    return [s for s in sld.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def timing(sld):
    return sld._element.find(f"{{{P}}}timing")


def clicks(sld) -> int:
    return len(timing(sld).findall(f".//{{{P}}}cTn[@nodeType='clickEffect']"))


# -- каталог библиотеки --------------------------------------------------------


def test_catalog_loads_and_matches_shapes(lib):
    found = pptx_templates.load(lib)

    assert list(found) == ["cover", "big_number", "question_choice", "steps"]
    assert found["question_choice"].slots[1].kind == "bullets"
    assert found["question_choice"].default_order == ["options", "answer"]
    # Сверка фигур (`slot:*`, `deco:*` на своём слайде) прошла внутри загрузки


def test_catalog_with_missing_slot_shape_is_refused(tmp_path):
    broken = build_library(tmp_path)
    text = (broken / "catalog.toml").read_text(encoding="utf-8")
    (broken / "catalog.toml").write_text(
        text.replace('name = "topic"', 'name = "absent"'), encoding="utf-8"
    )

    with pytest.raises(ValueError, match="slot:absent"):
        pptx_templates.load(broken)


def test_catalog_without_library_is_unavailable(tmp_path):
    assert pptx_templates.available(tmp_path / "nowhere") is False


def test_index_is_compact_and_names_limits(lib):
    text = pptx_templates.index(pptx_templates.load(lib))

    assert "[middle,end]" in text
    assert "options (bullets, до 4, до 40 зн., обязательно)" in text
    assert "caption (text, до 60 зн.)" in text


# -- схема сценария --------------------------------------------------------------


def test_every_deck_field_has_a_description():
    """По описаниям режим правки объясняет модели, что как править: поле без
    описания для неё не существует — включая записи слотов и добавленного."""
    schema = strict_schema(PresentationDeckDraft)
    slide_schema = schema["properties"]["slides_middle"]["items"]

    for name, prop in slide_schema["properties"].items():
        assert prop.get("description", "").strip(), name
    for entry in ("fill", "add"):
        items_schema = slide_schema["properties"][entry]["items"]
        for sub_name, sub in items_schema["properties"].items():
            assert sub.get("description", "").strip(), (entry, sub_name)


def test_deck_schema_is_strict_for_the_provider():
    schema = strict_schema(PresentationDeckDraft)

    slide_schema = schema["properties"]["slides_middle"]["items"]
    assert slide_schema["additionalProperties"] is False
    assert "anyOf" not in schema


def test_any_form_accepts_both_slide_shapes():
    draft = any_form_draft()
    old = {"layout": "title", "title": "Т", "notes": "н"}
    new = dict(scenario_slide("cover"), fill=[fill("topic", text="Тема")])
    content = {
        "goals": ["Научимся"],
        "slides_beginning": [new],
        "slides_middle": [old],
        "slides_end": [dict(scenario_slide("steps"), fill=[fill("steps", items=["Раз"])])],
    }

    checked = draft.model_validate(content)

    assert checked.slides_beginning[0].template == "cover"
    assert checked.slides_middle[0].layout == "title"
    assert checked.slides_end[0].template == "steps"


# -- сборка колоды ----------------------------------------------------------------


def test_deck_builds_and_reopens(lib):
    data, opened = built(lib, question_slide(), scenario_slide("big_number"))

    assert len(opened.slides) == 4  # титульный, цели и два слайда сценария
    assert data[:2] == b"PK"
    # Исходные слайды библиотеки из колоды вынуты: демо-текста нигде нет
    assert all("Тема урока" not in texts(s) for s in opened.slides)


def test_slots_are_filled(lib):
    _, opened = built(lib, question_slide())

    sld = opened.slides[2]
    assert "Where is my bag?" in texts(sld)
    options = shapes_by_name(sld)["slot:options"]
    assert [p.text for p in options.text_frame.paragraphs] == [
        "It is here",
        "It is under the chair",
    ]


def test_deco_is_removed_only_when_named_and_only_deco(lib):
    _, removed = built(lib, scenario_slide("cover", remove=["deco:border"]))
    _, kept = built(lib, scenario_slide("cover"))

    assert "deco:border" not in shapes_by_name(removed.slides[2])
    assert "deco:border" in shapes_by_name(kept.slides[2])
    # Слот украшением не считается: чужое имя не удаляется и ругается
    out = deck.render_deck(
        deck_content(scenario_slide("cover", remove=["slot:topic"])), HEADER, lib
    )
    assert any("не украшение" in w for w in out.warnings)


# -- пустые слоты и украшения-спутники -------------------------------------------------


def test_unfilled_optional_slot_is_removed_with_placeholder(lib):
    """Пустой необязательный слот уходит целиком, с текстом-заполнителем
    шаблона: в комплекте учителя заполнителей нет (спека, раздел 3, п. 4)."""
    _, opened = built(lib, scenario_slide("big_number", fill=[fill("number", text="42")]))

    sld = opened.slides[2]
    assert "slot:caption" not in shapes_by_name(sld)
    assert "Подпись" not in texts(sld)


def test_empty_fill_counts_as_unfilled(lib):
    out = deck.render_deck(
        deck_content(
            scenario_slide(
                "question_choice",
                fill=[fill("question", text="Где?"), fill("options", items=["А"]), fill("answer")],
            )
        ),
        HEADER,
        lib,
    )
    opened = Presentation(BytesIO(out.data))

    assert "slot:answer" not in shapes_by_name(opened.slides[2])
    assert not any("«answer»" in w for w in out.warnings)  # слот необязательный


def test_deco_following_empty_slot_is_removed(lib):
    """Украшение-спутник из deco.follows уходит вместе с опустевшим слотом,
    само, без `remove` в сценарии."""
    _, without = built(
        lib,
        scenario_slide(
            "question_choice",
            fill=[fill("question", text="Где?"), fill("options", items=["А", "Б"])],
        ),
    )
    _, with_picture = built(lib, question_slide())

    assert "deco:picture_bg" not in shapes_by_name(without.slides[2])
    assert "deco:picture_bg" in shapes_by_name(with_picture.slides[2])


def test_deco_group_is_removed_whole(lib):
    """deco: может быть группой фигур — снимается целиком, вместе с детьми."""
    _, opened = built(lib, scenario_slide("cover", remove=["deco:border"]))
    _, kept = built(lib, scenario_slide("cover"))

    assert "deco:border" not in shapes_by_name(opened.slides[2])
    children = {
        part.name
        for shape in kept.slides[2].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP
        for part in shape.shapes
    }
    assert "border-part-1" in children and "border-part-2" in children


def test_follows_pointing_at_unknown_slot_is_refused(tmp_path):
    broken = build_library(tmp_path)
    text = (broken / "catalog.toml").read_text(encoding="utf-8")
    (broken / "catalog.toml").write_text(
        text.replace('follows = { picture_bg = "picture" }', 'follows = { picture_bg = "absent" }'),
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="follows"):
        pptx_templates.load(broken)


def test_default_animation_skips_unfilled_slots_silently(lib):
    out = deck.render_deck(
        deck_content(
            scenario_slide(
                "question_choice",
                fill=[fill("question", text="Где?"), fill("options", items=["А"])],
            )
        ),
        HEADER,
        lib,
    )

    assert not any("animation.order" in w for w in out.warnings)


# -- списки и кегль -------------------------------------------------------------------


def test_bullets_keep_template_numbering(lib):
    """Каждый пункт наследует абзац-образец шаблона: автонумерация та же,
    что в шаблоне (у вариантов — буквы A. B. C., у шагов — 1. 2. 3.),
    и кегль не ниже границы (спека, п. 6)."""
    _, opened = built(
        lib,
        question_slide(),
        scenario_slide("steps", fill=[fill("title", text="Как"), fill("steps", items=["Раз"])]),
    )

    options = shapes_by_name(opened.slides[2])["slot:options"]
    for paragraph in options.text_frame.paragraphs:
        num = paragraph._p.pPr.find(qn("a:buAutoNum"))
        assert num is not None and num.get("type") == "alphaUcPeriod"
        assert paragraph.runs[0].font.size.pt >= 24
    steps = shapes_by_name(opened.slides[3])["slot:steps"]
    num = steps.text_frame.paragraphs[0]._p.pPr.find(qn("a:buAutoNum"))
    assert num is not None and num.get("type") == "arabicPeriod"
    assert steps.text_frame.paragraphs[0].runs[0].font.size.pt >= 24


def test_preview_shows_template_markers(lib):
    data, _ = built(
        lib,
        question_slide(),
        scenario_slide("steps", fill=[fill("title", text="Как"), fill("steps", items=["Раз"])]),
    )

    html = lesson_preview.render("pptx", data, "ru")

    assert ">A. </span>" in html  # буквы шаблона, а не «1. 2. 3.»
    assert ">1. </span>" in html


def test_library_picture_travels_with_the_slide(lib):
    """Клон слайда переезжает вместе со связями: картинка библиотеки —
    та же, байт в байт."""
    _, opened = built(lib, scenario_slide("steps"))

    stamps = pictures(opened.slides[2])
    assert len(stamps) == 1
    assert stamps[0].image.blob == (images.FILES / "check.png").read_bytes()


def test_slot_image_fits_frame_and_placeholder_is_gone(lib):
    frame = shapes_by_name(Presentation(str(lib / "library.pptx")).slides[2])["slot:picture"]
    _, opened = built(lib, question_slide())

    sld = opened.slides[2]
    # Картинка заменила рамку и переняла её имя slot:<name>
    named = [s for s in sld.shapes if s.name == "slot:picture"]
    assert len(named) == 1 and named[0].shape_type == MSO_SHAPE_TYPE.PICTURE
    pic = named[0]
    # Вписана в рамку слота с сохранением пропорций (PNG каталога квадратный)
    assert pic.left >= frame.left and pic.top >= frame.top
    assert pic.left + pic.width <= frame.left + frame.width + 1000
    assert pic.height / pic.width == pytest.approx(1.0, abs=0.02)


def test_unknown_image_key_leaves_slot_empty_with_warning(lib):
    out = deck.render_deck(
        deck_content(
            scenario_slide(
                "question_choice",
                fill=[
                    fill("question", text="Где?"),
                    fill("options", items=["А", "Б"]),
                    fill("picture", image="not_in_catalog"),
                ],
            )
        ),
        HEADER,
        lib,
    )

    assert any("not_in_catalog" in w for w in out.warnings)
    opened = Presentation(BytesIO(out.data))
    # Слот пустой, а не картинка с чужим ключом
    assert "slot:picture" in shapes_by_name(opened.slides[2])
    assert pictures(opened.slides[2]) == []


def test_add_items_land_at_percent_coordinates(lib):
    added = {"kind": "shape", "x": 50, "y": 50, "w": 10, "h": 10, "style": "circle"}
    hint = {
        "kind": "text",
        "x": 5,
        "y": 88,
        "w": 40,
        "h": 8,
        "text": "Подсказка: is",
        "style": "caption",
    }
    _, opened = built(lib, scenario_slide("cover", add=[hint, added]))

    sld = opened.slides[2]
    ovals = [
        s
        for s in sld.shapes
        if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and s._element.spPr.find(qn("a:prstGeom")).get("prst") == "ellipse"
    ]
    assert len(ovals) == 1
    assert ovals[0].left == pytest.approx(12192000 * 0.5, rel=0.01)
    box = next(s for s in sld.shapes if s.has_text_frame and "Подсказка" in s.text_frame.text)
    assert box.text_frame.paragraphs[0].runs[0].font.size.pt >= 18


# -- анимация ----------------------------------------------------------------------


def test_animation_by_order_and_catalog_default(lib):
    """`order` модели главнее умолчания каталога; пустой порядок берёт
    умолчание; none — слайд без эффектов вовсе."""
    _, by_order = built(
        lib, question_slide(animation={"effect": "fade", "order": ["answer", "options"]})
    )
    _, by_default = built(lib, question_slide())
    _, without = built(lib, question_slide(animation={"effect": "none", "order": []}))

    assert clicks(by_order.slides[2]) == 2
    assert clicks(by_default.slides[2]) == 2  # умолчание каталога: options, answer
    assert 'filter="fade"' in etree.tostring(timing(by_order.slides[2])).decode()
    assert timing(without.slides[2]) is None
    assert by_order.slides[2]._element.find(f"{{{P}}}transition") is not None
    assert without.slides[2]._element.find(f"{{{P}}}transition") is None


def test_animation_target_is_the_inserted_picture(lib):
    """Слот-картинка открывается как рисунок, а не как оставленный
    под ним заполнитель."""
    _, opened = built(
        lib, question_slide(animation={"effect": "appear", "order": ["picture"]})
    )

    target = timing(opened.slides[2]).find(f".//{{{P}}}spTgt")
    assert int(target.get("spid")) in [p.shape_id for p in pictures(opened.slides[2])]


def test_unknown_order_name_warns(lib):
    out = deck.render_deck(
        deck_content(question_slide(animation={"effect": "appear", "order": ["ghost"]})),
        HEADER,
        lib,
    )

    assert any("ghost" in w for w in out.warnings)


def test_timing_is_structurally_valid(lib):
    """Номера узлов уникальны, цели эффектов — существующие фигуры
    (ECMA-376 19.5, как у старой сборки)."""
    _, opened = built(
        lib,
        question_slide(),
        scenario_slide("big_number", fill=[fill("number", text="42")]),
    )

    for number in (2, 3):
        timing_el = timing(opened.slides[number])
        if timing_el is None:
            continue
        ids = [n.get("id") for n in timing_el.iter(f"{{{P}}}cTn")]
        assert len(ids) == len(set(ids)) and ids[0] == "1"
        shape_ids = {s.shape_id for s in opened.slides[number].shapes}
        for tgt in timing_el.iter(f"{{{P}}}spTgt"):
            assert int(tgt.get("spid")) in shape_ids


# -- заметки и предупреждения проверки ------------------------------------------------


def test_notes_are_set(lib):
    _, opened = built(lib, question_slide(notes="Спросить двух учеников"))

    assert opened.slides[2].notes_slide.notes_text_frame.text == "Спросить двух учеников"


def test_unknown_template_is_skipped_with_warning(lib):
    out = deck.render_deck(
        deck_content(scenario_slide("no_such_template", notes="н")), HEADER, lib
    )

    assert any("no_such_template" in w for w in out.warnings)
    # Титульный и цели на месте, слайд-пустышка не добавлен
    assert len(Presentation(BytesIO(out.data)).slides) == 2


def test_required_slot_left_empty_warns(lib):
    out = deck.render_deck(
        deck_content(scenario_slide("big_number", fill=[fill("caption", text="подпись")])),
        HEADER,
        lib,
    )

    assert any("обязательный слот «number»" in w for w in out.warnings)


def test_old_slide_builds_into_the_same_deck(lib):
    old = {"layout": "image_text", "title": "Вспомним", "lines": ["40 + 7"], "notes": "Шаг 1"}
    data = deck.render_deck(deck_content(question_slide(), old), HEADER, lib).data
    opened = Presentation(BytesIO(data))

    assert "40 + 7" in texts(opened.slides[3])


# -- переполнение ----------------------------------------------------------------


def test_overflow_shrinks_to_the_floor_and_warns(lib):
    long = "слово " * 30
    out = deck.render_deck(
        deck_content(scenario_slide("cover", fill=[fill("topic", text=long)])),
        HEADER,
        lib,
    )
    opened = Presentation(BytesIO(out.data))
    topic = shapes_by_name(opened.slides[2])["slot:topic"]

    assert topic.text_frame.paragraphs[0].runs[0].font.size.pt == 24  # нижняя граница слотов
    assert any("длиннее лимита" in w for w in out.warnings)


def test_text_within_limit_keeps_template_size(lib):
    _, opened = built(lib, scenario_slide("cover", fill=[fill("topic", text="Короткая тема")]))

    topic = shapes_by_name(opened.slides[2])["slot:topic"]
    assert topic.text_frame.paragraphs[0].runs[0].font.size.pt == 44  # кегль шаблона


def test_caption_add_shrinks_to_18(lib):
    hint = {
        "kind": "text",
        "x": 5,
        "y": 88,
        "w": 20,
        "h": 5,
        "text": "подпись " * 20,
        "style": "caption",
    }
    _, opened = built(lib, scenario_slide("cover", add=[hint]))

    box = next(
        s
        for s in opened.slides[2].shapes
        if s.has_text_frame and s.text_frame.text.startswith("подпись")
    )
    assert box.text_frame.paragraphs[0].runs[0].font.size.pt == 18


def test_number_longer_than_twelve_chars_warns(lib):
    out = deck.render_deck(
        deck_content(
            scenario_slide("big_number", fill=[fill("number", text="1 + 2 + 3 + 4 + 5")])
        ),
        HEADER,
        lib,
    )

    assert any("длиннее лимита (12 зн.)" in w for w in out.warnings)


# -- выбор пути по форме содержимого -------------------------------------------------


def test_scenario_content_goes_through_deck_path(lib, monkeypatch):
    monkeypatch.setattr(pptx_templates, "root", lambda: lib)
    data = lesson_files.render("presentation", deck_content(question_slide()), HEADER)
    opened = Presentation(BytesIO(data))

    assert "Where is my bag?" in texts(opened.slides[2])


def test_old_content_goes_the_old_way(lib, monkeypatch):
    """Материалы прежней формы собираются старым путём даже там, где
    библиотека уже есть: определяем по форме содержимого, не по дате."""
    monkeypatch.setattr(pptx_templates, "root", lambda: lib)
    old = {
        "goals": ["Научимся"],
        "slides_beginning": [
            {"layout": "image_text", "title": "Вспомним", "lines": ["40 + 7"], "notes": "Шаг 1"}
        ],
        "slides_middle": [],
        "slides_end": [],
    }
    data = lesson_files.render("presentation", old, HEADER)
    opened = Presentation(BytesIO(data))

    assert "40 + 7" in texts(opened.slides[2])


def test_scenario_without_library_falls_back_to_the_old_path(monkeypatch):
    """Откат деплоя без библиотеки при сценарных материалах: выгрузка
    живёт, слайд упрощается, ошибка — в лог."""
    monkeypatch.setattr(pptx_templates, "root", lambda: Path("nowhere"))
    monkeypatch.setattr(
        "app.adapters.lesson_files.log", logging.getLogger("test-deck-no-library")
    )

    data = lesson_files.render("presentation", deck_content(question_slide()), HEADER)
    opened = Presentation(BytesIO(data))

    assert len(opened.slides) == 3  # титульный, цели и пустой слайд


def test_edit_validation_accepts_scenario_slide():
    item = question_slide()

    checked = BlockRewriter(ai=None).validate("presentation", "slides_middle", [item])

    assert checked[0]["template"] == "question_choice"
    assert edit_tools._item(item)["value"] == item


def test_outline_shows_template_and_main_text():
    lines = kit_outline(
        [
            {
                "kind": "presentation",
                "status": "ready",
                "blocks": [
                    {
                        "name": "slides_middle",
                        "server": False,
                        "content": [
                            question_slide(),
                            scenario_slide("cover", fill=[fill("topic", text="Старая форма")]),
                        ],
                    }
                ],
            }
        ]
    )
    text = "\n".join(lines)

    assert "слайд 3 [question_choice] «Where is my bag?» img:backpack" in text
    assert "слайд 4 [cover] «Старая форма»" in text


def test_deck_prompt_carries_catalog_rules_and_images(lib):
    deck_system = prompt.deck_system(pptx_templates.index(pptx_templates.load(lib)))

    for key in pptx_templates.load(lib):
        assert key in deck_system
    assert "два раза подряд" in deck_system  # правило чередования шаблонов
    assert all(key in deck_system for key in images.catalog())
    assert "{images}" not in deck_system


def test_classic_prompt_keeps_image_keys():
    assert all(key in prompt.SYSTEM for key in images.catalog())
    assert "{images}" not in prompt.SYSTEM


def test_edit_chat_prompt_carries_the_template_catalog(lib, monkeypatch):
    monkeypatch.setattr(pptx_templates, "root", lambda: lib)

    block = edit_chat.templates_block()

    assert "question_choice" in block
    assert "прежней формы" in block


def test_edit_chat_prompt_without_library_has_no_catalog(monkeypatch, tmp_path):
    """Настоящая библиотека лежит по пути по умолчанию — отсутствие каталога
    в промпте проверяется на заведомо несуществующем пути настройки."""
    monkeypatch.setattr(pptx_templates, "root", lambda: tmp_path / "nowhere")

    assert edit_chat.templates_block() == ""


# -- предпросмотр ------------------------------------------------------------------


def test_preview_draws_template_slides(lib):
    data, opened = built(lib, question_slide())

    html = lesson_preview.render("pptx", data, "ru")

    assert html.count('class="slide"') == len(opened.slides)
    assert "data:image/png;base64," in html
    text = re.sub(r"<[^>]+>", " ", html)
    for needle in ("Where is my bag?", "It is here", "It is under the chair"):
        assert needle in text, needle
    assert "<script" not in html
