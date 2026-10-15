"""Маленькая библиотека шаблонов для тестов (спека, раздел 1).

Настоящей библиотеки в репозитории ещё нет — её делает дизайн параллельно,
и тесты исполнителя не должны её ждать. Здесь код собирает `library.pptx`
на четыре шаблонных слайда и пишет `catalog.toml` рядом: фигуры названы
по спеке (`slot:*`, `deco:*`), на одном слайде — настоящая картинка, чтобы
проверить, что клон слайда переезжает вместе со связями. Кегли и рамки —
рабочие, а не красивые: на них проверяется вписывание текста и картинок.

Собирается в фикстуре тестов во временный каталог; коммитить бинарный
`.pptx` незачем — источник читаемее файла.
"""

import tomllib
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from app.lesson_ai import images

# 16:9, как у настоящей библиотеки
WIDTH, HEIGHT = Emu(12192000), Emu(6858000)

CATALOG = """\
# Библиотека-фикстура для тестов исполнителя: формат — спека, раздел 1.
# Слайды идут в порядке catalog.toml, номера — с единицы.

[[template]]
key = "cover"
slide = 1
title = "Титул"
purpose = "Тема урока крупно и подзаголовок"
stage = ["beginning"]

  [[template.slot]]
  name = "topic"
  kind = "text"
  max_chars = 60
  required = true

  [template.deco]
  removable = ["border"]

[[template]]
key = "big_number"
slide = 2
title = "Крупное число"
purpose = "Число или слово крупно и подпись под ним"
stage = ["middle", "end"]

  [[template.slot]]
  name = "number"
  kind = "number"
  required = true

  [[template.slot]]
  name = "caption"
  kind = "text"
  max_chars = 60

[[template]]
key = "question_choice"
slide = 3
title = "Вопрос с вариантами"
purpose = "Проверка понимания: вопрос, варианты, ответ по клику"
stage = ["middle", "end"]

  [[template.slot]]
  name = "question"
  kind = "text"
  max_chars = 90
  required = true

  [[template.slot]]
  name = "options"
  kind = "bullets"
  max_items = 4
  max_chars = 40
  required = true

  [[template.slot]]
  name = "answer"
  kind = "text"
  max_chars = 60

  [[template.slot]]
  name = "picture"
  kind = "image"

  [template.deco]
  removable = ["picture_bg"]
  follows = { picture_bg = "picture" }

  [template.animation]
  default = ["options", "answer"]

[[template]]
key = "steps"
slide = 4
title = "Шаги"
purpose = "Нумерованные шаги алгоритма"
stage = ["middle"]

  [[template.slot]]
  name = "title"
  kind = "text"
  max_chars = 40
  required = true

  [[template.slot]]
  name = "steps"
  kind = "bullets"
  max_items = 5
  max_chars = 60
  required = true

  [template.deco]
  removable = ["wave"]
"""


def build(directory: Path) -> Path:
    """Собрать библиотеку в `directory`; вернуть путь к ней."""
    deck = Presentation()
    deck.slide_width, deck.slide_height = WIDTH, HEIGHT
    blank = deck.slide_layouts[6]

    cover = deck.slides.add_slide(blank)
    _text(cover, "slot:topic", 0.9, 2.0, 11.5, 2.2, 44, "Тема урока")
    # Рамка-украшение как группа: декорация может быть и группой фигур
    border = cover.shapes.add_group_shape()
    _deco(border.shapes, "border-part-1", 0.4, 0.4, 12.5, 0.25)
    _deco(border.shapes, "border-part-2", 0.4, 6.85, 12.5, 0.25)
    border.name = "deco:border"

    fact = deck.slides.add_slide(blank)
    _text(fact, "slot:number", 0.9, 1.4, 11.5, 2.8, 96, "42")
    _text(fact, "slot:caption", 0.9, 4.4, 11.5, 1.4, 28, "Подпись")

    question = deck.slides.add_slide(blank)
    _text(question, "slot:question", 0.9, 0.6, 11.5, 1.5, 36, "Вопрос?")
    _bullets(question, "slot:options", 0.9, 2.4, 7.0, 3.4, 28, "Вариант", "alphaUcPeriod")
    _text(question, "slot:answer", 0.9, 6.0, 11.5, 1.0, 32, "Ответ")
    _deco(question.shapes, "deco:picture_bg", 8.5, 2.3, 4.0, 3.6)
    _frame(question, "slot:picture", 8.6, 2.4, 3.8, 3.4)

    steps = deck.slides.add_slide(blank)
    _text(steps, "slot:title", 0.9, 0.5, 11.5, 1.2, 40, "Шаги")
    _bullets(steps, "slot:steps", 0.9, 2.0, 11.5, 4.6, 28, "Шаг", "arabicPeriod")
    _deco(steps.shapes, "deco:wave", 0.0, 6.9, 13.34, 0.6)
    # Настоящая картинка со связью: клон слайда обязан перевезти и её
    blob = (images.FILES / "check.png").read_bytes()
    picture = steps.shapes.add_picture(BytesIO(blob), Inches(10.2), Inches(0.4), Inches(1.4))
    picture.name = "art:stamp"

    directory.mkdir(parents=True, exist_ok=True)
    deck.save(directory / "library.pptx")
    (directory / "catalog.toml").write_text(CATALOG, encoding="utf-8")
    return directory


def read_catalog(directory: Path) -> dict:
    return tomllib.loads((directory / "catalog.toml").read_text(encoding="utf-8"))


def _text(slide, name: str, x: float, y: float, w: float, h: float, pt: int, text: str):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    frame = box.text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(pt)
    run.font.name = "Arial"
    return box


def _deco(shapes, name: str, x: float, y: float, w: float, h: float):
    """Украшение; первый аргумент — коллекция фигур слайда или группы."""
    shape = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    return shape


def _bullets(
    slide, name: str, x: float, y: float, w: float, h: float, pt: int, text: str, num_type: str
):
    """Слот-список с автонумерацией шаблона: исполнитель обязан сохранять
    её в каждом пункте (`buAutoNum` в `pPr` абзаца-образца)."""
    box = _text(slide, name, x, y, w, h, pt, text)
    ppr = box.text_frame.paragraphs[0]._p.get_or_add_pPr()
    num = ppr.makeelement(qn("a:buAutoNum"), {"type": num_type})
    ppr.append(num)
    return box


def _frame(slide, name: str, x: float, y: float, w: float, h: float):
    """Рамка слота-картинки: прямоугольник без текста, как в настоящей
    библиотеке (спека: «прямоугольник или заполнитель картинки»)."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    return shape
