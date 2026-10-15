"""«Жоспар»: просмотр материала на сайте — `GET /lesson_ai/materials/{id}/preview`.

Просмотр обещает учителю то же, что он распечатает. Поэтому тесты сверяют
страницу с выданным файлом, а не с содержимым материала: каждый абзац
документа и каждая строка слайда и заметок есть в тексте страницы. Плюс то,
что нельзя сломать: текст модели и учителя экранируется, склеенная ячейка
этапа остаётся одной, отказы — те же, что у файла.
"""

import re
from html.parser import HTMLParser
from io import BytesIO

import docx
import pytest
from pptx import Presentation

from app.adapters import lesson_preview
from tests.conftest import login
from tests.test_lesson_ai_docx import (
    NOT_FOUND,
    NOT_READY,
    OTHER_PHONE,
    all_lines,
    ready_material,
)
from tests.test_lesson_ai_files import (
    KINDS,
    full_content,
    kinds_for,
    material_id,
    ready_kit,
    rendered,
)
from tests.test_lesson_ai_kits import Scene, order

pytestmark = pytest.mark.usefixtures("lesson_ai_open")


@pytest.fixture
def scene() -> Scene:
    return Scene()


@pytest.fixture
def ru_objective_id(client, sms) -> int:
    login(client, sms)
    return Scene().objective.id


class _Text(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self.tags: list[str] = []

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        # Блочные теги — граница слов, иначе соседние ячейки слипнутся
        self.parts.append(" ")

    def handle_data(self, data):
        self.parts.append(data)


def page_text(html: str) -> str:
    parser = _Text()
    parser.feed(html)
    return squeeze("".join(parser.parts))


def page_tags(html: str) -> set[str]:
    parser = _Text()
    parser.feed(html)
    return set(parser.tags)


def squeeze(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def preview(client, material: int):
    return client.get(f"/lesson_ai/materials/{material}/preview")


def file_lines(kind: str, data: bytes) -> list[str]:
    """Весь текст файла: абзацы документа или строки слайдов с заметками."""
    if kind != "presentation":
        return all_lines(docx.Document(BytesIO(data)))
    lines = []
    for slide in Presentation(BytesIO(data)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines += [p.text for p in shape.text_frame.paragraphs]
        lines += slide.notes_slide.notes_text_frame.text.splitlines()
    return lines


@pytest.mark.parametrize("kind", KINDS)
def test_preview_shows_every_line_of_the_file(client, ai, ru_objective_id, kind):
    kit = ready_kit(client, ai, ru_objective_id, kinds_for(kind))
    material = material_id(kit, kind)
    data = client.get(f"/lesson_ai/materials/{material}/file").content

    resp = preview(client, material)

    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("text/html")
    assert "default-src 'none'" in resp.headers["content-security-policy"]
    text = page_text(resp.text)
    missing = [line for line in file_lines(kind, data) if squeeze(line) not in text]
    assert not missing, missing[:5]
    assert "script" not in page_tags(resp.text)


def test_presentation_preview_has_a_frame_per_slide(client, ai, ru_objective_id):
    kit = ready_kit(client, ai, ru_objective_id, ["presentation"])
    material = material_id(kit, "presentation")
    data = client.get(f"/lesson_ai/materials/{material}/file").content
    slides = len(Presentation(BytesIO(data)).slides)

    html = preview(client, material).text

    assert html.count('class="slide"') == slides
    assert f"1 / {slides}" in html


def test_lesson_plan_stage_cell_stays_merged(client, sms, ai, scene):
    """Этап урока в Word — одна ячейка на несколько шагов; в просмотре
    она тоже одна (`rowspan`), а не пустые клетки под названием этапа."""
    ids = ready_material(client, sms, ai, scene)

    html = preview(client, ids["material_id"]).text

    assert re.search(r'rowspan="[2-9]"', html)


def test_text_from_model_is_escaped():
    """В содержимом — ответ модели и правка учителя: разметка из них
    попадает на страницу текстом, а не тегами."""
    evil = '<script>alert("x")</script><img src=x onerror=alert(1)>'
    homework = full_content("homework") | {"intro": evil}
    slides = full_content("presentation") | {"goals": [evil]}
    for kind, fmt, content in (("homework", "docx", homework), ("presentation", "pptx", slides)):
        html = lesson_preview.render(fmt, rendered(kind, content), "ru")

        assert "script" not in page_tags(html), kind
        # Картинки презентации — свои, data:-адресом; чужая разметка тегом не становится
        assert all(
            'src="data:' in tag and "onerror" not in tag for tag in re.findall(r"<img[^>]*>", html)
        )
        assert evil in page_text(html), kind


def test_notes_label_is_in_the_language_of_the_kit():
    data = rendered("presentation", full_content("presentation", "kz"), "kz")

    html = lesson_preview.render("pptx", data, "kz")

    assert "Ескертпелер" in html
    assert "Заметки" not in html
    assert '<html lang="kk">' in html


# -- отказы — те же, что у файла ----------------------------------------------


def test_preview_requires_login(client, sms, ai, scene):
    ids = ready_material(client, sms, ai, scene)
    client.cookies.clear()

    assert preview(client, ids["material_id"]).status_code == 401


def test_someone_elses_preview_is_not_found(client, sms, ai, scene):
    ids = ready_material(client, sms, ai, scene)
    login(client, sms, phone=OTHER_PHONE)

    resp = preview(client, ids["material_id"])

    assert resp.status_code == 404, resp.text
    assert resp.json() == NOT_FOUND


def test_queued_material_has_no_preview(client, sms, scene):
    login(client, sms)
    [material] = order(client, scene.objective.id).json()["materials"]

    resp = preview(client, material["id"])

    assert resp.status_code == 409, resp.text
    assert resp.json() == NOT_READY
