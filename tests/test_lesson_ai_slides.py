"""«Жоспар»: презентация — типы слайдов, картинки из каталога, анимации PowerPoint.

Без базы: всё это чистые функции над содержимым материала — схема слайда,
каталог картинок, сборка `.pptx` и его просмотр на сайте.
"""

import logging
import re
from io import BytesIO
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.adapters import lesson_preview
from app.adapters.docx.common import DocHeader
from app.adapters.pptx.presentation import render_presentation
from app.lesson_ai import edit_tools, images
from app.lesson_ai.blocks import BlockRewriter
from app.lesson_ai.kit_map import kinds_catalog, kit_outline
from app.lesson_ai.prompts import presentation as prompt
from app.lesson_ai.schemas import strict_schema
from app.lesson_ai.schemas.presentation import (
    ANIMATIONS,
    DEFAULT_ANIMATION,
    LAYOUTS,
    PresentationDraft,
    Slide,
)

HEADER = DocHeader(lang="ru", teacher_name="Нурланова Айгуль", grade=3, topic="Глагол to be")
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

SLIDES = {
    "title": {"title": "Game time!", "lines": ["Find your pair"], "image": "dice"},
    "agenda": {"title": "Сегодня", "lines": ["Узнаем am, is, are", "Поиграем"]},
    "image_text": {
        "title": "Who is this?",
        "lines": ["This is a cat.", "It is small."],
        "image": "cat",
    },
    "two_columns": {
        "title": "am / is / are",
        "left_title": "One",
        "lines": ["I am", "He is"],
        "right_title": "Many",
        "right_lines": ["We are", "They are"],
    },
    "big_fact": {"title": "Remember", "fact": "I am", "lines": ["I am Ayan."], "image": "boy"},
    "question": {
        "title": "She ___ a teacher.",
        "lines": ["am", "is", "are"],
        "answer": "is",
        "image": "teacher",
    },
    "steps": {"title": "Steps", "lines": ["Look", "Choose", "Say"], "image": "magnifier"},
    "summary": {"title": "Today", "lines": ["I am", "He is"], "image": "star"},
}


def slide(layout: str, **over) -> dict:
    return {
        "layout": layout,
        "animation": DEFAULT_ANIMATION[layout],
        "notes": "Заметка",
        **SLIDES[layout],
        **over,
    }


def deck(*slides: dict) -> dict:
    return {
        "goals": ["Научимся"],
        "slides_beginning": list(slides[:1]),
        "slides_middle": list(slides[1:]),
        "slides_end": [],
    }


def built(*slides: dict):
    data = render_presentation(deck(*slides), HEADER)
    return data, Presentation(BytesIO(data))


def timing(sld) -> etree._Element | None:
    return sld._element.find(f"{{{P}}}timing")


# -- схема ---------------------------------------------------------------------


def test_every_slide_field_has_a_description():
    """По описаниям режим правки объясняет модели, что как править: поле без
    описания для неё не существует."""
    schema = strict_schema(PresentationDraft)
    slide_schema = schema["properties"]["slides_middle"]["items"]
    assert set(slide_schema["properties"]) == set(Slide.model_fields)
    assert slide_schema["required"] == list(Slide.model_fields)
    for name, prop in slide_schema["properties"].items():
        assert prop.get("description", "").strip(), name


def test_schema_offers_all_layouts_and_animations():
    props = strict_schema(PresentationDraft)["properties"]["slides_middle"]["items"]["properties"]
    assert props["layout"]["enum"] == list(LAYOUTS) and len(LAYOUTS) == 8
    assert props["animation"]["enum"] == list(ANIMATIONS) == ["none", "appear", "fade"]
    assert props["image"]["type"] == "string" and "enum" not in props["image"]
    # Все типы названы в описании типа: по нему помощник знает, какие бывают
    assert all(layout in props["layout"]["description"] for layout in LAYOUTS)


def test_old_slide_without_new_fields_passes_validation():
    old = {"title": "Вспомним", "lines": ["40 + 7"], "visual": "палочки", "notes": "Шаг 1"}

    new = Slide.model_validate(old)

    assert (new.layout, new.image, new.animation) == ("image_text", "", "appear")
    # Правка из чата проверяет блок целиком той же схемой: старые соседи проходят
    checked = BlockRewriter(ai=None).validate(
        "presentation", "slides_middle", [old, slide("question")]
    )
    assert checked[0]["layout"] == "image_text" and checked[1]["answer"] == "is"


def test_default_animation_follows_layout():
    title = Slide.model_validate({"layout": "title", "title": "Т", "notes": "н"})
    fact = Slide.model_validate({"layout": "big_fact", "title": "Т", "notes": "н"})
    assert (title.animation, fact.animation) == ("none", "fade")


def test_unknown_image_key_is_dropped_with_a_log_without_the_key(caplog):
    with caplog.at_level(logging.WARNING):
        checked = Slide.model_validate({"title": "Т", "notes": "н", "image": "secret_key_42"})

    assert checked.image == ""
    assert any("каталог" in r.message for r in caplog.records)
    assert "secret_key_42" not in caplog.text


def test_known_image_key_is_kept_in_any_case():
    assert Slide.model_validate({"title": "Т", "notes": "н", "image": " Cat "}).image == "cat"


def test_new_slide_is_accepted_by_edit_tools_validation():
    item = slide("two_columns")
    checked = BlockRewriter(ai=None).validate("presentation", "slides_middle", [item])
    assert checked[0]["right_lines"] == ["We are", "They are"]
    assert edit_tools._item(item)["value"] == item


# -- каталог картинок ---------------------------------------------------------


def test_catalog_loads_and_every_file_is_a_png():
    found = images.catalog()
    assert 150 <= len(found) <= 500
    total = 0
    for image in found.values():
        data = (images.FILES / image.file).read_bytes()
        assert data[:8] == b"\x89PNG\r\n\x1a\n", image.key
        total += len(data)
        assert image.ru and image.kz and image.en, image.key
    assert total < 5_000_000
    # Ни одного лишнего файла: каталог и папка расходиться не должны
    assert {i.file for i in found.values()} == {p.name for p in images.FILES.glob("*.png")}


def test_catalog_has_attribution():
    text = (Path(images.ROOT) / "LICENSE.md").read_text(encoding="utf-8")
    assert "MIT" in text and "Microsoft" in text and "fluentui-emoji" in text


def test_search_finds_by_ru_kz_and_en_words():
    assert images.search("яблоко")[0].key == "apple"
    assert images.search("мысық")[0].key == "cat"
    assert images.search("rocket")[0].key == "rocket"
    assert images.search("абвгдежз") == []


def test_prompt_lists_every_key_in_the_system_prompt():
    assert all(key in prompt.SYSTEM for key in images.catalog())
    assert "{images}" not in prompt.SYSTEM


def test_find_image_tool():
    out = edit_tools.EditTools().call("find_image", {"query": "кошка"}, ctx=None)
    assert out["images"][0]["key"] == "cat"
    assert edit_tools.EditTools().call("find_image", {"query": " "}, ctx=None)["error"]


# -- сборка .pptx --------------------------------------------------------------


@pytest.mark.parametrize("layout", LAYOUTS)
def test_every_layout_builds_and_reopens(layout):
    data, opened = built(slide("image_text"), slide(layout))

    assert len(opened.slides) == 4  # титульный, цели и два слайда
    texts = [
        p.text for s in opened.slides[3].shapes if s.has_text_frame for p in s.text_frame.paragraphs
    ]
    assert SLIDES[layout]["title"] in texts or layout == "big_fact"
    assert data[:2] == b"PK"


def test_old_slide_without_layout_builds_like_before():
    old = {
        "title": "Вспомним",
        "lines": ["40 + 7", "56 − 6"],
        "visual": "палочки",
        "notes": "Шаг 1",
    }

    _, opened = built(old)

    lines = [
        p.text for s in opened.slides[2].shapes if s.has_text_frame for p in s.text_frame.paragraphs
    ]
    assert "Вспомним" in lines and "40 + 7" in lines and "56 − 6" in lines
    assert "Что показать: палочки" in opened.slides[2].notes_slide.notes_text_frame.text
    assert timing(opened.slides[2]) is not None  # умолчание старого слайда — appear


def pictures(sld) -> list:
    return [s for s in sld.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_picture_is_in_the_file_and_unknown_key_gives_no_picture():
    _, opened = built(slide("image_text", image="cat"), slide("image_text", image="not_in_catalog"))

    assert len(pictures(opened.slides[2])) == 1
    assert pictures(opened.slides[2])[0].image.content_type == "image/png"
    assert pictures(opened.slides[3]) == []


def test_slide_without_image_gets_full_width_text():
    _, with_image = built(slide("image_text", image="cat"))
    _, without = built(slide("image_text", image=""))

    def widest(sld):
        plate = next(s for s in sld.shapes if s.has_text_frame and "This is" in s.text_frame.text)
        return plate.width

    assert widest(without.slides[2]) > widest(with_image.slides[2])


def test_animated_slides_have_timing_and_transition_in_schema_order():
    _, opened = built(
        slide("image_text"), slide("question"), slide("big_fact"), slide("title", animation="none")
    )

    for number in (2, 3, 4):
        children = [etree.QName(c).localname for c in opened.slides[number]._element]
        assert children[-3:] == ["clrMapOvr", "transition", "timing"], children
    assert timing(opened.slides[5]) is None  # none — без эффектов
    assert opened.slides[5]._element.find(f"{{{P}}}transition") is None
    assert timing(opened.slides[0]) is None  # титульный из комплекта


def test_timing_is_structurally_valid():
    """Проверка по ECMA-376 19.5 без PowerPoint: уникальные номера узлов, цели
    эффектов — существующие фигуры (и абзацы), каждая цель с текстом в bldLst."""
    _, opened = built(
        slide("image_text"),
        slide("two_columns", animation="fade"),
        slide("steps"),
        slide("question"),
        slide("big_fact"),
    )
    ns = {"p": P}
    for sld in list(opened.slides)[1:]:
        timing_el = timing(sld)
        if timing_el is None:
            continue
        ids = [n.get("id") for n in timing_el.iter(f"{{{P}}}cTn")]
        assert len(ids) == len(set(ids)) and ids[0] == "1"
        shapes = {s.shape_id: s for s in sld.shapes}
        built_ids = {b.get("spid") for b in timing_el.findall(".//p:bldP", ns)}
        for tgt in timing_el.iter(f"{{{P}}}spTgt"):
            spid = int(tgt.get("spid"))
            assert spid in shapes and tgt.get("spid") in built_ids
            for rng in tgt.iter(f"{{{P}}}pRg"):
                assert int(rng.get("end")) < len(shapes[spid].text_frame.paragraphs)
        for node in timing_el.iter(f"{{{P}}}cTn"):
            if node.get("nodeType") in ("clickEffect", "withEffect"):
                assert node.get("presetClass") == "entr" and node.get("grpId") == "0"
        # Каждый щелчок ждёт клика
        clicks = timing_el.findall(".//p:cTn[@nodeType='mainSeq']/p:childTnLst/p:par/p:cTn", ns)
        assert clicks and all(
            c.find("p:stCondLst/p:cond", ns).get("delay") == "indefinite" for c in clicks
        )


def test_list_items_appear_one_click_each_and_question_answer_on_click():
    _, opened = built(slide("image_text"), slide("question"), slide("steps"))

    def clicks(sld) -> int:
        return len(timing(sld).findall(f".//{{{P}}}cTn[@nodeType='clickEffect']"))

    assert clicks(opened.slides[2]) == 2  # две строки
    assert clicks(opened.slides[3]) == 1  # ответ
    assert clicks(opened.slides[4]) == 3  # три шага (плашка и кружок — один щелчок)
    steps = timing(opened.slides[4])
    assert len(steps.findall(f".//{{{P}}}cTn[@nodeType='withEffect']")) == 3


def test_fade_adds_fade_effect_and_appear_does_not():
    _, opened = built(
        slide("image_text", animation="fade"), slide("image_text", animation="appear")
    )

    assert 'filter="fade"' in etree.tostring(timing(opened.slides[2])).decode()
    assert "animEffect" not in etree.tostring(timing(opened.slides[3])).decode()


def test_question_answer_is_not_a_part_of_the_visible_title():
    _, opened = built(slide("question"))
    sld = opened.slides[2]
    answer = next(
        s for s in sld.shapes if s.has_text_frame and s.text_frame.text.startswith("Ответ")
    )
    target = timing(sld).find(f".//{{{P}}}spTgt")
    assert int(target.get("spid")) == answer.shape_id


def test_text_never_below_readable_size():
    _, opened = built(*(slide(layout) for layout in LAYOUTS))
    sizes = [
        run.font.size.pt
        for sld in opened.slides
        for s in sld.shapes
        if s.has_text_frame
        for p in s.text_frame.paragraphs
        for run in p.runs
        # метка этапа — для учителя, а не с доски
        if not run.text.endswith("урока")
    ]
    assert min(sizes) >= 22


def test_kazakh_deck_keeps_language_and_labels():
    header = DocHeader(lang="kz", teacher_name="Т", grade=3, topic="Тақырып")
    data = render_presentation(deck(slide("image_text"), slide("question")), header)
    langs = set(Presentation(BytesIO(data)).slides[3]._element.xpath("//a:rPr/@lang"))
    assert langs == {"kk-KZ"}
    texts = [
        s.text_frame.text for s in Presentation(BytesIO(data)).slides[3].shapes if s.has_text_frame
    ]
    assert any(t.startswith("Жауап") for t in texts)


# -- просмотр на сайте -----------------------------------------------------------


def test_preview_draws_new_layouts_and_pictures():
    data, opened = built(*(slide(layout) for layout in LAYOUTS))

    html = lesson_preview.render("pptx", data, "ru")

    assert html.count('class="slide"') == len(opened.slides)
    assert "data:image/png;base64," in html
    assert re.search(r"border-radius:50%", html)  # кружки
    text = re.sub(r"<[^>]+>", " ", html)
    for needle in ("Ответ: is", "I am", "We are", "Many", "1. ", "She ___ a teacher."):
        assert needle in text, needle
    assert "<script" not in html


def test_preview_picture_is_data_uri_only():
    data, _ = built(slide("image_text", image="cat"))
    html = lesson_preview.render("pptx", data, "ru")
    assert not re.search(r'src="(?!data:)', html)


def test_outline_shows_layout_and_image():
    lines = kit_outline(
        [
            {
                "kind": "presentation",
                "status": "ready",
                "blocks": [
                    {
                        "name": "slides_middle",
                        "server": False,
                        "content": [
                            slide("question", title="Where is my bag?", image="backpack"),
                            {"title": "Old"},
                        ],
                    }
                ],
            }
        ]
    )
    text = "\n".join(lines)
    assert "[question] «Where is my bag?» img:backpack" in text
    assert "слайд 4 «Old»" in text


def test_edit_catalog_describes_new_slide_fields():
    text = kinds_catalog()
    for field in ("layout", "image", "animation", "answer", "fact", "right_lines"):
        assert f"- {field}:" in text
    assert "find_image" in text
