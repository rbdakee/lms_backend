"""Библиотека шаблонов слайдов сходится со своим каталогом (PPTX_TEMPLATES_SPEC.md,
раздел 1): исполнитель сценария ищет фигуры по именам из каталога, и слот без
фигуры или фигура без записи всплыли бы только в собранном комплекте учителя."""

import tomllib
from pathlib import Path

import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent / "app" / "lesson_ai" / "pptx_templates"
KINDS = {"text", "bullets", "number", "image"}
STAGES = {"beginning", "middle", "end"}


@pytest.fixture(scope="module")
def catalog():
    with open(ROOT / "catalog.toml", "rb") as f:
        return tomllib.load(f)["template"]


@pytest.fixture(scope="module")
def slides():
    return list(Presentation(str(ROOT / "library.pptx")).slides)


def names(slide):
    return [shape.name for shape in slide.shapes]


def test_one_slide_per_entry_in_order(catalog, slides):
    assert 12 <= len(catalog) <= 15
    assert len(slides) == len(catalog)
    assert [t["slide"] for t in catalog] == list(range(1, len(catalog) + 1))
    keys = [t["key"] for t in catalog]
    assert len(set(keys)) == len(keys)


def test_entries_are_complete(catalog):
    for t in catalog:
        assert t["key"].isascii() and t["title"] and t["purpose"], t["key"]
        assert t["stage"] and set(t["stage"]) <= STAGES, t["key"]
        assert t["slot"], t["key"]
        for s in t["slot"]:
            assert s["kind"] in KINDS, (t["key"], s["name"])
            assert isinstance(s["required"], bool), (t["key"], s["name"])
            if s["kind"] != "image":
                assert s["max_chars"] > 0, (t["key"], s["name"])
            if s["kind"] == "number":
                assert s["max_chars"] <= 12, (t["key"], s["name"])
            if s["kind"] == "bullets":
                assert s["max_items"] > 0, (t["key"], s["name"])


def test_every_slot_has_its_shape_and_no_stray_slots(catalog, slides):
    for t, slide in zip(catalog, slides, strict=True):
        shapes = [n for n in names(slide) if n.startswith("slot:")]
        assert len(shapes) == len(set(shapes)), t["key"]
        assert sorted(shapes) == sorted(f"slot:{s['name']}" for s in t["slot"]), t["key"]


def test_removable_and_animation_point_at_real_shapes(catalog, slides):
    for t, slide in zip(catalog, slides, strict=True):
        decos = {n[len("deco:") :] for n in names(slide) if n.startswith("deco:")}
        assert set(t["deco"]["removable"]) == decos, t["key"]
        slots = {s["name"] for s in t["slot"]}
        assert set(t["animation"]["default"]) <= slots, t["key"]


def test_follows_point_at_real_slots(catalog):
    """Украшение-спутник следует за настоящими слотами шаблона; число пунктов
    («слот:N») — только у bullets и в пределах max_items."""
    for t in catalog:
        slots = {s["name"]: s for s in t["slot"]}
        for deco, value in (t["deco"].get("follows") or {}).items():
            conditions = value if isinstance(value, list) else [value]
            assert deco in t["deco"]["removable"], (t["key"], deco)
            for condition in conditions:
                name, _, count = str(condition).partition(":")
                assert name in slots, (t["key"], deco, condition)
                if count:
                    slot = slots[name]
                    assert slot["kind"] == "bullets" and int(count) <= slot["max_items"], (
                        t["key"],
                        condition,
                    )


def test_text_slots_are_large_enough_for_the_board(catalog, slides):
    # Смарт-доска с пяти метров: текст в слотах не мельче 24 pt
    for t, slide in zip(catalog, slides, strict=True):
        kinds = {f"slot:{s['name']}": s["kind"] for s in t["slot"]}
        for shape in slide.shapes:
            if kinds.get(shape.name, "image") == "image":
                continue
            sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs]
            assert sizes and min(sizes) >= 24, (t["key"], shape.name)


def test_previews_exist(catalog):
    for t in catalog:
        assert (ROOT / "previews" / f"{t['key']}.png").is_file(), t["key"]
