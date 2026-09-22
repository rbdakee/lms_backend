"""«Жоспар»: шесть новых видов файлом — презентация в PowerPoint и пять
документов Word (`GET /lesson_ai/materials/{id}/file`).

Как и у ҚМЖ (`test_lesson_ai_docx.py`), проверяется выданный файл, а не код
рендера: он открывается тем же python-pptx или python-docx, и в нём ищется
текст. Главное, что нельзя сломать: ответов нет ни на одном листе, который
уходит детям, — они только в ключах (TEACHER_AI_BRIEF, раздел 4.2); цель
обучения с кодом — в шапке каждого документа, а у презентации — в заметках,
не на доске; язык подписей — язык обучения (раздел 1); призы игры рисует
сервер в каждом файле (раздел 4.3; L5_games_db_design, раздел 1.4).

Ответы модели — образцы из `tests/samples/lesson_ai/`, генерация настоящая,
ИИ подставной, задания очереди берутся руками (`test_lesson_ai_kits.py`).
"""

import copy
import json
import re
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote

import docx
import pytest
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from pptx import Presentation

from app.adapters import lesson_files
from app.adapters.docx import answer_key, assessment, game, homework, tasks
from app.adapters.docx.common import DocHeader
from app.adapters.pptx import presentation
from app.domain.lesson_ai import ANSWER_SOURCES, assemble_answer_key
from tests.conftest import login
from tests.test_lesson_ai_kazakh import (
    KZ_OBJECTIVE,
    KZ_SECTION,
    KZ_SUBSECTION,
    KZ_TOPIC,
    kz_objective,
)
from tests.test_lesson_ai_kits import CODE, OBJECTIVE_TEXT, TOPIC, Scene, run_next

# Доступ к «Жоспару» (L7) в этих тестах открыт: они проверяют продукт,
# а сама проверка доступа — в `test_lesson_ai_access.py`
pytestmark = pytest.mark.usefixtures("lesson_ai_open")

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
SAMPLES = Path(__file__).parent / "samples" / "lesson_ai"

MODEL_KINDS = ("presentation", "tasks", "homework", "assessment", "game")
DOCX_KINDS = ("tasks", "homework", "assessment", "game", "answer_key")
KINDS = (*MODEL_KINDS, "answer_key")

# Имена файлов — своей копией, а не из `lesson_files.TITLES`: тест, берущий
# ожидание из того же словаря, что и код, промолчит на русском слове
# в казахском имени
TITLES = {
    "ru": {
        "presentation": "Презентация",
        "tasks": "Задания",
        "homework": "Домашнее задание",
        "assessment": "Проверочная работа",
        "game": "Игра",
        "answer_key": "Ключи с ответами",
    },
    "kz": {
        "presentation": "Презентация",
        "tasks": "Тапсырмалар",
        "homework": "Үй тапсырмасы",
        "assessment": "Формативті бағалау",
        "game": "Ойын",
        "answer_key": "Жауаптар кілті",
    },
}

# Буквы уровня — кириллица, как на карточке ребёнка: латинская «A» выглядит
# так же, но в поиске по документу и в проверке орфографии это другая буква
LEVELS = {"cards_support": "А", "cards_core": "В", "cards_challenge": "С"}
LEVEL_LABEL = {"ru": "Уровень {letter}", "kz": "{letter} деңгейі"}

SERVER = {
    "ru": {
        "section": {"title": "Числа и величины", "subsection": "Операции над числами"},
        "learning_objectives": [{"code": CODE, "marker": "", "text": OBJECTIVE_TEXT}],
    },
    "kz": {
        "section": {"title": KZ_SECTION, "subsection": KZ_SUBSECTION},
        "learning_objectives": [{"code": CODE, "marker": "", "text": KZ_OBJECTIVE}],
    },
}
OBJECTIVE_LINE = {"ru": f"{CODE} {OBJECTIVE_TEXT}", "kz": f"{CODE} {KZ_OBJECTIVE}"}
HEADERS = {
    "ru": DocHeader(lang="ru", teacher_name="Нурланова Айгуль", grade=2, topic=TOPIC),
    "kz": DocHeader(lang="kz", teacher_name="Сәрсенқызы Әйгерім", grade=2, topic=KZ_TOPIC),
}


def sample(kind: str) -> dict:
    """Ответ модели по схеме вида — свежей копией: тесты его правят."""
    return json.loads((SAMPLES / f"{kind}.json").read_text(encoding="utf-8"))


def key_content(lang: str = "ru") -> dict:
    """Содержимое ключей, как его собирает сервер из трёх образцов."""
    sources = [(kind, list(sample(kind).items())) for kind in ANSWER_SOURCES]
    return SERVER[lang] | assemble_answer_key(sources)


def full_content(kind: str, lang: str = "ru") -> dict:
    """Содержимое готового материала: ответ модели и блоки сервера."""
    if kind == "answer_key":
        return key_content(lang)
    return SERVER[lang] | sample(kind)


def rendered(kind: str, content: dict, lang: str = "ru") -> bytes:
    return lesson_files.render(kind, content, HEADERS[lang])


def strings(value) -> list[str]:
    """Все строки содержимого, на любой глубине."""
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [item for part in value.values() for item in strings(part)]
    if isinstance(value, list):
        return [item for part in value for item in strings(part)]
    return []


def without_answers(value):
    if isinstance(value, dict):
        return {key: without_answers(item) for key, item in value.items() if key != "answer"}
    if isinstance(value, list):
        return [without_answers(item) for item in value]
    return value


def answers(value) -> list[str]:
    if isinstance(value, dict):
        own = [value["answer"]] if isinstance(value.get("answer"), str) else []
        return own + [item for part in value.values() for item in answers(part)]
    if isinstance(value, list):
        return [item for part in value for item in answers(part)]
    return []


def distinctive_answers(content: dict) -> list[str]:
    """Ответы, которых нет больше нигде в материале. Короткий ответ «17»
    честно стоит в условии «57 − 17», и по нему утечку не отличить; ответ
    с ходом решения в условии не встречается — его появление в листе
    ребёнка и есть утечка."""
    rest = "\n".join(strings(without_answers(content)))
    return [answer for answer in answers(content) if answer not in rest]


# -- чтение файлов ---------------------------------------------------------


def body_lines(data: bytes) -> list[str]:
    """Текст документа абзацами в порядке тела — и вне таблиц, и в ячейках.
    Перенос строки внутри абзаца (`w:br`) остаётся «\\n»."""
    body = docx.Document(BytesIO(data)).element.body
    return [Paragraph(p, None).text for p in body.iter(qn("w:p"))]


def body_text(data: bytes) -> str:
    return "\n".join(body_lines(data))


def slides_of(data: bytes) -> list:
    return list(Presentation(BytesIO(data)).slides)


def slide_lines(slide) -> list[str]:
    return [
        paragraph.text
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        if paragraph.text
    ]


def notes_of(slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def runs_of(data: bytes) -> list:
    return [
        run
        for slide in slides_of(data)
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ]


# Порядок детей tblPr и tcPr задан схемой OOXML (последовательность, а не
# выбор): Word на нарушенном порядке предлагает «восстановить» файл
TBL_PR_ORDER = [
    qn(f"w:{tag}")
    for tag in (
        "tblStyle",
        "tblpPr",
        "tblOverlap",
        "bidiVisual",
        "tblStyleRowBandSize",
        "tblStyleColBandSize",
        "tblW",
        "jc",
        "tblCellSpacing",
        "tblInd",
        "tblBorders",
        "shd",
        "tblLayout",
        "tblCellMar",
        "tblLook",
    )
]
TC_PR_ORDER = [
    qn(f"w:{tag}")
    for tag in (
        "cnfStyle",
        "tcW",
        "gridSpan",
        "hMerge",
        "vMerge",
        "tcBorders",
        "shd",
        "noWrap",
        "tcMar",
        "textDirection",
        "tcFitText",
        "vAlign",
        "hideMark",
    )
]


def assert_word_opens_cleanly(data: bytes) -> None:
    """Рамки таблиц — в таблице и на своём месте в tblPr, свойства ячеек —
    в порядке схемы: иначе Word открывает файл через «восстановить»."""
    body = docx.Document(BytesIO(data)).element.body
    for tbl in body.iter(qn("w:tbl")):
        tags = [child.tag for child in tbl.tblPr]
        assert tags.count(qn("w:tblBorders")) <= 1
        known = [tag for tag in tags if tag in TBL_PR_ORDER]
        assert known == sorted(known, key=TBL_PR_ORDER.index), tags
    for tc_pr in body.iter(qn("w:tcPr")):
        tags = [child.tag for child in tc_pr]
        known = [tag for tag in tags if tag in TC_PR_ORDER]
        assert known == sorted(known, key=TC_PR_ORDER.index), tags
        assert len(known) == len(set(known)), tags


# -- комплект через ручку ----------------------------------------------------


def ready_kit(client, ai, objective_id: int, kinds, topic: str = TOPIC) -> dict:
    """Заказ и генерация всех видов по образцам. Очередь берёт задания
    в порядке видов брифа, ключи — последними: их ответ не нужен."""
    resp = client.post(
        "/lesson_ai/kits", json={"objective_id": objective_id, "topic": topic, "kinds": kinds}
    )
    assert resp.status_code == 201, resp.text
    for kind in [kind for kind in MODEL_KINDS if kind in kinds]:
        ai.returns_json(sample(kind))
    for _kind in kinds:
        assert run_next(ai)
    kit = client.get(f"/lesson_ai/kits/{resp.json()['id']}").json()
    assert kit["status"] == "ready", kit
    return kit


def kinds_for(kind: str) -> list[str]:
    # Ключам нужен источник, иначе заказ не примут
    return ["tasks", "homework", "assessment", "answer_key"] if kind == "answer_key" else [kind]


def material_id(kit: dict, kind: str) -> int:
    [material] = [m for m in kit["materials"] if m["kind"] == kind]
    return material["id"]


def download(client, kit: dict, kind: str):
    resp = client.get(f"/lesson_ai/materials/{material_id(kit, kind)}/file")
    assert resp.status_code == 200, resp.text
    return resp


@pytest.fixture
def ru_objective_id(client, sms) -> int:
    login(client, sms)
    return Scene().objective.id


@pytest.fixture
def kz_objective_id(client, sms) -> int:
    objective = kz_objective()
    login(client, sms)
    return objective.id


# -- ручка: тип и имя файла -------------------------------------------------


@pytest.mark.parametrize("kind", KINDS)
def test_file_type_and_name_follow_the_kind(client, ai, ru_objective_id, kind):
    """Учитель складывает материалы урока в одну папку: имя — «Вид — тема»
    на языке обучения в `filename*`, запасное из ASCII — в `filename`.
    Презентация — PowerPoint, остальное — Word."""
    ext = "pptx" if kind == "presentation" else "docx"
    kit = ready_kit(client, ai, ru_objective_id, kinds_for(kind))

    resp = download(client, kit, kind)

    assert resp.headers["content-type"] == (PPTX if kind == "presentation" else DOCX)
    disposition = resp.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    fallback = re.search(r'filename="([^"]*)"', disposition).group(1)
    assert fallback == f"{kind}_{kit['id']}.{ext}"
    encoded = re.search(r"filename\*=UTF-8''(\S+)", disposition).group(1)
    assert encoded.isascii()
    assert unquote(encoded) == f"{TITLES['ru'][kind]} — {TOPIC}.{ext}"
    # Байты — настоящий файл своего формата
    if ext == "pptx":
        assert slides_of(resp.content)
    else:
        assert body_lines(resp.content)


@pytest.mark.parametrize("kind", KINDS)
def test_kazakh_kit_gives_kazakh_file_names(client, ai, kz_objective_id, kind):
    """Казахский комплект — казахское имя файла: «Тапсырмалар», «Ойын»,
    «Жауаптар кілті». Язык — от цели, а не от интерфейса."""
    ext = "pptx" if kind == "presentation" else "docx"
    kit = ready_kit(client, ai, kz_objective_id, kinds_for(kind), topic=KZ_TOPIC)
    assert kit["lang"] == "kz"

    resp = download(client, kit, kind)

    disposition = resp.headers["content-disposition"]
    assert re.search(r'filename="([^"]*)"', disposition).group(1) == f"{kind}_{kit['id']}.{ext}"
    encoded = re.search(r"filename\*=UTF-8''(\S+)", disposition).group(1)
    assert unquote(encoded) == f"{TITLES['kz'][kind]} — {KZ_TOPIC}.{ext}"


# -- презентация ------------------------------------------------------------


def test_presentation_has_title_goals_and_every_slide(client, ai, ru_objective_id):
    """Слайдов ровно: титульный, «Сегодня на уроке» и все слайды трёх
    блоков по порядку. На титульном — тема и класс из комплекта."""
    kit = ready_kit(client, ai, ru_objective_id, ["presentation"])
    content = sample("presentation")

    slides = slides_of(download(client, kit, "presentation").content)

    stages = [
        slide
        for block in ("slides_beginning", "slides_middle", "slides_end")
        for slide in content[block]
    ]
    assert len(slides) == 2 + len(stages) == 6
    assert slide_lines(slides[0]) == [TOPIC, "2 класс"]
    assert slide_lines(slides[1]) == ["Сегодня на уроке", *content["goals"]]
    for slide, item in zip(slides[2:], stages, strict=True):
        assert item["title"] in slide_lines(slide)
        assert all(line in slide_lines(slide) for line in item["lines"])


def test_objective_code_is_in_title_notes_only(client, ai, ru_objective_id):
    """Цель с кодом — для учителя, а не для детей: в заметках титульного
    слайда, и ни на одном слайде её не видно с доски."""
    kit = ready_kit(client, ai, ru_objective_id, ["presentation"])

    slides = slides_of(download(client, kit, "presentation").content)

    assert OBJECTIVE_LINE["ru"] in notes_of(slides[0])
    for slide in slides:
        on_screen = "\n".join(slide_lines(slide))
        assert CODE not in on_screen
        assert OBJECTIVE_TEXT not in on_screen


def test_visual_and_notes_go_to_their_slide_notes(client, ai, ru_objective_id):
    """«Что показать» и слова учителя — в заметках своего слайда: картинок
    в файле нет, а подсказка к слайду, уехавшая на соседний, путает."""
    kit = ready_kit(client, ai, ru_objective_id, ["presentation"])
    content = sample("presentation")
    stages = [
        item
        for block in ("slides_beginning", "slides_middle", "slides_end")
        for item in content[block]
    ]

    slides = slides_of(download(client, kit, "presentation").content)

    for slide, item in zip(slides[2:], stages, strict=True):
        notes = notes_of(slide)
        assert item["notes"] in notes
        assert f"Что показать: {item['visual']}" in notes
        on_screen = "\n".join(slide_lines(slide))
        assert item["notes"] not in on_screen
        assert item["visual"] not in on_screen


def test_presentation_font_is_arial_everywhere(client, ai, ru_objective_id):
    kit = ready_kit(client, ai, ru_objective_id, ["presentation"])

    runs = runs_of(download(client, kit, "presentation").content)

    assert runs
    assert {run.font.name for run in runs} == {"Arial"}
    assert {run.font._rPr.get("lang") for run in runs} == {"ru-RU"}


def test_kazakh_presentation_is_kazakh(client, ai, kz_objective_id):
    """Казахский класс — казахские подписи слайдов и язык проверки
    орфографии у каждого текста, иначе PowerPoint подчеркнёт всё."""
    kit = ready_kit(client, ai, kz_objective_id, ["presentation"], topic=KZ_TOPIC)

    data = download(client, kit, "presentation").content

    slides = slides_of(data)
    assert slide_lines(slides[0]) == [KZ_TOPIC, "2 сынып"]
    assert slide_lines(slides[1])[0] == "Бүгін сабақта"
    assert OBJECTIVE_LINE["kz"] in notes_of(slides[0])
    assert notes_of(slides[0]).startswith("Оқу мақсаты:")
    lines = [line for slide in slides for line in slide_lines(slide)]
    labels = [line for line in lines if line in ("Сабақтың басы", "Сабақтың ортасы")]
    assert labels
    for russian in presentation.LABELS["ru"].values():
        assert russian not in lines
    # Заметки образца русские («Начало урока, шаг 1») — это содержимое;
    # подписи рендера в заметках — казахские
    notes = "\n".join(notes_of(slide) for slide in slides)
    assert "Не көрсету керек: " in notes
    assert "Что показать" not in notes
    assert "Цель обучения" not in notes
    runs = runs_of(data)
    assert {run.font._rPr.get("lang") for run in runs} == {"kk-KZ"}
    assert {run.font.name for run in runs} == {"Arial"}


# -- документы Word: ответы только в ключах ---------------------------------


@pytest.mark.parametrize("kind", ANSWER_SOURCES)
def test_student_sheets_carry_no_answers(client, ai, ru_objective_id, kind):
    """Лист уходит детям — ответов на нём нет ни одного: они в ключах,
    отдельным файлом у учителя (TEACHER_AI_BRIEF, раздел 4.2)."""
    kit = ready_kit(client, ai, ru_objective_id, [kind])
    leaks = distinctive_answers(sample(kind))
    # Без этой сверки проверка ниже молча проверяла бы ничего
    assert len(leaks) >= 4, leaks

    text = body_text(download(client, kit, kind).content)

    assert [answer for answer in leaks if answer in text] == []


@pytest.mark.parametrize("lang", ["ru", "kz"])
@pytest.mark.parametrize("kind", DOCX_KINDS)
def test_objective_with_code_heads_every_document(kind, lang):
    """Цель обучения с кодом — в шапке каждого документа, до первого
    задания: завуч сверяет с ней лист, не листая его."""
    content = full_content(kind, lang)

    lines = body_lines(rendered(kind, content, lang))

    line = OBJECTIVE_LINE[lang]
    found = [number for number, text in enumerate(lines) if line in text]
    assert found
    first_task = min(
        (
            number
            for number, text in enumerate(lines)
            for task in answers_free_texts(content)
            if task in text
        ),
        default=len(lines),
    )
    assert found[0] < first_task


def answers_free_texts(content: dict) -> list[str]:
    """Условия задач материала — то, что должно идти после шапки."""
    texts = []

    def walk(value):
        if isinstance(value, dict):
            if isinstance(value.get("text"), str) and "answer" in value:
                texts.append(value["text"].split("\n")[0])
            if isinstance(value.get("task"), str):
                texts.append(value["task"].split("\n")[0])
            for item in value.values():
                walk(item)
        elif isinstance(value, list):
            for item in value:
                walk(item)

    walk(content)
    return [text for text in texts if len(text) > 12]


def test_assessment_repeats_objective_on_each_sheet():
    """Каждый вариант — свой лист со своей шапкой, и таблица критериев
    тоже: учитель печатает страницы по отдельности."""
    lines = body_lines(rendered("assessment", full_content("assessment")))

    assert sum(OBJECTIVE_LINE["ru"] in line for line in lines) == 3


# -- язык подписей ------------------------------------------------------------


RENDER_MODULES = {
    "tasks": tasks,
    "homework": homework,
    "assessment": assessment,
    "game": game,
    "answer_key": answer_key,
}


def label_strings(labels: dict) -> list[str]:
    """Подписи словаря; у шаблонов — постоянная часть до подстановки."""
    out = []
    for value in labels.values():
        if isinstance(value, dict):
            out += label_strings(value)
        else:
            out += [part.strip() for part in re.split(r"\{[a-z_]+\}", value) if part.strip()]
    return out


@pytest.mark.parametrize("kind", DOCX_KINDS)
def test_kazakh_document_has_no_russian_labels(kind):
    """Язык документа — язык обучения. Содержимое образцов русское, поэтому
    из текста сначала вычищается всё, что пришло из материала: остаток —
    подписи рендера, и русской среди них быть не должно."""
    content = full_content(kind, "kz")
    text = body_text(rendered(kind, content, "kz"))
    pieces = strings(content) + [line for s in strings(content) for line in s.split("\n")]
    for piece in sorted(set(pieces), key=len, reverse=True):
        if len(piece) > 2:
            text = text.replace(piece, "")

    module = RENDER_MODULES[kind]
    kazakh = set(label_strings(module.LABELS["kz"]))
    russian = [
        label
        for label in label_strings(module.LABELS["ru"])
        if label not in kazakh and len(label) > 2
    ]
    assert russian
    assert [label for label in russian if label in text] == []
    # И казахские подписи на месте — хотя бы заголовок
    assert label_strings(module.LABELS["kz"])[0] in text


def test_kazakh_document_language_is_kk(client, ai, kz_objective_id):
    """Язык проверки орфографии в Word — казахский: иначе документ
    подчёркнут целиком как «русский с ошибками»."""
    kit = ready_kit(client, ai, kz_objective_id, kinds_for("answer_key"), topic=KZ_TOPIC)

    for kind in ANSWER_SOURCES + ("answer_key",):
        document = docx.Document(BytesIO(download(client, kit, kind).content))
        rpr = document.styles["Normal"].element.rPr
        assert [node.get(qn("w:val")) for node in rpr.findall(qn("w:lang"))] == ["kk-KZ"]
        fonts = rpr.find(qn("w:rFonts"))
        assert fonts.get(qn("w:ascii")) == fonts.get(qn("w:cs")) == "Times New Roman"


# -- игра: призы рисует сервер -------------------------------------------------


def is_die_net(table) -> bool:
    cells = [cell.text for row in table.rows for cell in row.cells]
    return len(table.rows) == 4 and sorted(text for text in cells if text) == list("123456")


def coin_cells(data: bytes) -> set[str]:
    document = docx.Document(BytesIO(data))
    return {
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.endswith(" тг")
    }


@pytest.mark.parametrize(
    ("lang", "diploma", "tokens"),
    [("ru", "ГРАМОТА", "Призы. Жетоны"), ("kz", "МАҚТАУ ҚАҒАЗЫ", "Сыйлықтар. Жетондар")],
)
@pytest.mark.parametrize("content", [{}, None], ids=["empty", "sample"])
def test_game_always_has_tokens_and_diploma(lang, diploma, tokens, content):
    """Награда в конце игры обещана брифом (раздел 4.3): жетоны и грамота —
    в каждом файле игры, даже если модель о призах не написала ни слова."""
    content = full_content("game", lang) if content is None else content

    data = rendered("game", content, lang)

    lines = body_lines(data)
    assert diploma in lines
    assert tokens in lines
    document = docx.Document(BytesIO(data))
    label = (content.get("reward") or {}).get("token_label") or "+1"
    grids = [
        table
        for table in document.tables
        if [cell.text for row in table.rows for cell in row.cells] == [label] * 40
    ]
    assert len(grids) == 1


@pytest.mark.parametrize("needs_die", [True, False])
def test_die_net_only_when_rules_need_a_die(needs_die):
    content = full_content("game")
    content["rules"]["needs_die"] = needs_die

    data = rendered("game", content)

    document = docx.Document(BytesIO(data))
    assert any(is_die_net(table) for table in document.tables) is needs_die
    assert ("Кубик" in body_lines(data)) is needs_die


@pytest.mark.parametrize(("coins", "cells"), [([], set()), ([1, 5, 10], {"1 тг", "5 тг", "10 тг"})])
def test_coins_only_when_rules_list_them(coins, cells):
    content = full_content("game")
    content["rules"]["coins"] = coins

    data = rendered("game", content)

    assert coin_cells(data) == cells
    assert ("Игровые деньги" in body_lines(data)) is bool(coins)


def test_downloaded_game_follows_sample_rules(client, ai, ru_objective_id):
    """Через ручку — то же: у образца кубика и денег нет, призы есть,
    ключ ведущего — в файле игры."""
    kit = ready_kit(client, ai, ru_objective_id, ["game"])
    content = sample("game")
    assert content["rules"]["needs_die"] is False
    assert content["rules"]["coins"] == []

    data = download(client, kit, "game").content

    lines = body_lines(data)
    assert "ГРАМОТА" in lines
    assert not any(is_die_net(table) for table in docx.Document(BytesIO(data)).tables)
    assert coin_cells(data) == set()
    assert "Ключ для ведущего" in lines
    assert content["title"] in lines


# -- ключи с ответами ---------------------------------------------------------


def expected_key_groups(lang: str = "ru") -> list[tuple[str, list[dict]]]:
    """Группы ключей по образцам, в порядке документа: подпись и задачи."""
    groups = []
    for block, letter in LEVELS.items():
        level = LEVEL_LABEL[lang].format(letter=letter)
        for card in sample("tasks")[block]:
            groups.append((f"{level} — {card['title']}", card["tasks"]))
    work = sample("homework")
    groups += [(work[name]["title"], work[name]["tasks"]) for name in ("required", "optional")]
    test = sample("assessment")
    groups += [(test[name]["title"], test[name]["tasks"]) for name in ("variant_1", "variant_2")]
    return groups


def test_answer_key_prints_every_answer_by_group(client, ai, ru_objective_id):
    """Ключи — все ответы заданий, домашнего и проверочной: по группам,
    с номером задания, как его видит ребёнок, и с уровнем карточки —
    иначе две карточки с одним названием не различить."""
    kit = ready_kit(client, ai, ru_objective_id, kinds_for("answer_key"))

    data = download(client, kit, "answer_key").content

    document = docx.Document(BytesIO(data))
    groups = expected_key_groups()
    captions = [caption for caption, _tasks in groups]
    assert [p.text for p in document.paragraphs if p.text in captions] == captions
    for heading in ("Задания на урок", "Домашнее задание", "Проверочная работа"):
        assert heading in [p.text for p in document.paragraphs]
    tables = document.tables[1:]
    assert len(tables) == len(groups)
    for table, (_caption, items) in zip(tables, groups, strict=True):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        assert rows[0] == ["№", "Задание", "Ответ"]
        assert rows[1:] == [
            [str(number), item["text"], item["answer"]]
            for number, item in enumerate(items, start=1)
        ]
    # Каждый ответ, которого нет на листах детей, — здесь
    text = body_text(data)
    for kind in ANSWER_SOURCES:
        assert all(answer in text for answer in distinctive_answers(sample(kind)))


def test_kazakh_answer_key_labels_card_levels_in_kazakh():
    data = rendered("answer_key", key_content("kz"), "kz")

    lines = body_lines(data)
    assert "А деңгейі — " + sample("tasks")["cards_support"][0]["title"] in lines
    assert not any(line.startswith("Уровень ") for line in lines)


# -- рендер не падает и Word открывает без «восстановить» ---------------------


MINIMAL = {
    "presentation": {"slides_middle": [{"title": "Считаем", "lines": ["1 + 1"]}]},
    "tasks": {"cards_core": [{"title": "К", "tasks": [{"text": "1 + 1 = ___", "answer": "2"}]}]},
    "homework": {"required": {"title": "О", "tasks": [{"text": "Реши", "answer": "2"}]}},
    "assessment": {
        "variant_1": {
            "title": "Вариант 1",
            "tasks": [
                {"text": "Реши", "answer": "2", "descriptors": [{"text": "решает", "points": 1}]}
            ],
        }
    },
    "game": {"title": "Игра", "rules": {"steps": ["Шаг"], "needs_die": True, "coins": [5]}},
    "answer_key": {
        "tasks": [
            {
                "block": "cards_core",
                "title": "К",
                "answers": [{"number": 1, "task": "1 + 1", "answer": "2"}],
            }
        ]
    },
}


@pytest.mark.parametrize("lang", ["ru", "kz"])
@pytest.mark.parametrize("fill", ["empty", "minimal", "sample"])
@pytest.mark.parametrize("kind", KINDS)
def test_render_survives_any_fill(kind, fill, lang):
    """Рендер не падает на неполном содержимом: один пропавший ключ после
    ручной правки — 500 вместо документа. Word-файлы при этом открываются
    без «восстановить»."""
    content = {
        "empty": {},
        "minimal": MINIMAL[kind],
        "sample": full_content(kind, lang),
    }[fill]

    data = rendered(kind, copy.deepcopy(content), lang)

    if kind == "presentation":
        assert slides_of(data)
    else:
        assert_word_opens_cleanly(data)
        assert "None" not in body_text(data)


def test_word_order_check_catches_misplaced_borders():
    """Проверка порядка сама умеет падать: рамки, переставленные в конец
    tblPr, как было до L4, — ошибка."""
    document = docx.Document(BytesIO(rendered("tasks", full_content("tasks"))))
    tbl_pr = document.tables[0]._tbl.tblPr
    tbl_pr.append(tbl_pr.find(qn("w:tblBorders")))
    out = BytesIO()
    document.save(out)

    with pytest.raises(AssertionError):
        assert_word_opens_cleanly(out.getvalue())


# -- файл из текущего содержимого ---------------------------------------------


def test_edited_block_shows_up_in_the_next_download(client, ai, ru_objective_id):
    """Файл собирается при каждом скачивании: правка блока видна сразу,
    без перегенерации и без сохранённой копии файла."""
    kit = ready_kit(client, ai, ru_objective_id, ["tasks"])
    old = sample("tasks")["intro"]
    assert old in body_lines(download(client, kit, "tasks").content)

    new = "Сегодня каждая верная карточка — шаг каравана к Таразу."
    resp = client.put(
        f"/lesson_ai/materials/{material_id(kit, 'tasks')}/blocks/intro", json={"content": new}
    )
    assert resp.status_code == 200, resp.text

    lines = body_lines(download(client, kit, "tasks").content)
    assert new in lines
    assert old not in lines


def test_edited_slides_show_up_in_the_next_presentation(client, ai, ru_objective_id):
    kit = ready_kit(client, ai, ru_objective_id, ["presentation"])
    goals = ["Научимся считать десятками", "Поиграем в караван"]

    resp = client.put(
        f"/lesson_ai/materials/{material_id(kit, 'presentation')}/blocks/goals",
        json={"content": goals},
    )
    assert resp.status_code == 200, resp.text

    slides = slides_of(download(client, kit, "presentation").content)
    assert slide_lines(slides[1]) == ["Сегодня на уроке", *goals]
