"""«Жоспар»: ҚМЖ файлом Word — `GET /lesson_ai/materials/{id}/file`.

Документ учитель несёт завучу, и проверяют его глазами, а не по JSON. Поэтому
тесты не сверяют байты и не читают код рендера, а открывают выданный .docx
тем же python-docx и ищут в нём текст: каждый из пятнадцати блоков формы — на
своём месте, цель обучения — дословно из справочника с кодом, ход урока —
таблицей со склеенной ячейкой этапа (CONTRACT.md, «Жоспар: файл и правка
по блокам»; журнал L4, записи 3–6).

Генерация — настоящая, ИИ — подставной, задания очереди берутся руками,
как в `test_lesson_ai_kits.py`.
"""

import copy
import re
from io import BytesIO
from urllib.parse import unquote

import docx
import pytest
from docx.oxml.ns import qn

from app.adapters.docx.lesson_plan import LABELS, PlanHeader, render_lesson_plan
from app.lesson_ai.schemas.lesson_plan import STAGE_MINUTES
from tests.conftest import login, login_named
from tests.test_lesson_ai_kits import (
    CODE,
    DRAFT,
    OBJECTIVE_TEXT,
    TOPIC,
    Scene,
    generated,
    hurry_retries,
    order,
    run_next,
)

# Доступ к «Жоспару» (L7) в этих тестах открыт: они проверяют продукт,
# а сама проверка доступа — в `test_lesson_ai_access.py`
pytestmark = pytest.mark.usefixtures("lesson_ai_open")

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OTHER_PHONE = "+7 (707) 555-01-02"
NOT_FOUND = {"error": {"code": "not_found", "message": "Материал не найден"}}
NOT_READY = {"error": {"code": "material_not_ready", "message": "Материал ещё не готов"}}
RU = LABELS["ru"]

# Черновик с несколькими шагами в середине урока и несколькими пунктами
# в списках: на одном шаге и одном пункте не видно ни склейки ячейки этапа,
# ни потерянного второго элемента. Тексты все разные — тогда по тексту видно,
# в какую строку документа уехал блок.
DRAFT_MANY = copy.deepcopy(DRAFT)
DRAFT_MANY["lesson_goals"].append("Учащиеся смогут объяснить ход вычислений соседу по парте.")
DRAFT_MANY["assessment_criteria"].append("Учащийся называет разряды в числе 46.")
DRAFT_MANY["reflection"].append("Хватило ли времени на работу в парах?")
DRAFT_MANY["stage_middle"] += [
    {
        "teacher": "Даёт парам карточки: 52 + 34, 67 – 25, 43 + 16.",
        "students": "Решают в парах и сверяют ответы с соседями.",
        "resources": ["карточки", "тетрадь"],
    },
    {
        "teacher": "Проводит физминутку «Десятки — приседаем, единицы — хлопаем».",
        "students": "Выполняют движения под счёт.",
        "resources": [],
    },
]

STEP_STAGES = tuple(STAGE_MINUTES)


def download(client, material_id: int):
    return client.get(f"/lesson_ai/materials/{material_id}/file")


def opened(resp) -> docx.document.Document:
    assert resp.status_code == 200, resp.text
    return docx.Document(BytesIO(resp.content))


def cell_lines(cell) -> list[str]:
    return [paragraph.text for paragraph in cell.paragraphs]


def all_lines(document) -> list[str]:
    """Весь текст документа абзацами — и вне таблиц, и в каждой ячейке."""
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines += cell_lines(cell)
    return lines


def head_value(document, kind: str) -> list[str]:
    """Строки ячейки-значения в шапке ҚМЖ — по подписи слева. Искать надо
    именно по подписи: текст в документе, но в чужой строке формы, — ошибка."""
    [table] = [t for t in document.tables if t.rows[0].cells[0].text == RU["section"]]
    rows = [row for row in table.rows if row.cells[0].text == RU[kind]]
    assert len(rows) == 1, f"строка «{RU[kind]}» не одна в шапке"
    return cell_lines(rows[0].cells[1])


def stages_table(document):
    [table] = [t for t in document.tables if t.rows[0].cells[0].text == RU["stage"]]
    return table


def stage_rows(document, stage: str) -> list:
    """Строки хода урока, принадлежащие этапу, — по склеенной ячейке этапа."""
    return [
        row
        for row in stages_table(document).rows[1:]
        if row.cells[0].paragraphs[0].text == RU[stage]
    ]


def closing_value(document, kind: str) -> list[str]:
    for table in document.tables:
        head = [cell.text for cell in table.rows[0].cells]
        if RU[kind] in head and len(table.rows) > 1:
            return cell_lines(table.rows[1].cells[head.index(RU[kind])])
        if head[0] == RU[kind]:
            return cell_lines(table.rows[0].cells[1])
    raise AssertionError(f"в документе нет блока «{RU[kind]}»")


def bullets(items: list[str]) -> list[str]:
    return [f"• {item}" for item in items]


def ready_material(client, sms, ai, scene, draft=DRAFT_MANY, **kw) -> dict:
    """Комплект с готовым ҚМЖ; учитель вошёл без профиля."""
    kit = generated(client, sms, ai, scene, draft=draft, **kw)
    [material] = kit["materials"]
    assert material["status"] == "ready", material
    return {"kit_id": kit["id"], "material_id": material["id"]}


def ordered_by_named(client, sms, ai, scene, draft=DRAFT_MANY, **profile) -> dict:
    """То же, но учитель сначала заполнил профиль: ФИО шапки берётся оттуда."""
    login_named(client, sms, **profile)
    kit_id = order(client, scene.objective.id, class_size=24).json()["id"]
    ai.returns_json(draft)
    assert run_next(ai)
    [material] = client.get(f"/lesson_ai/kits/{kit_id}").json()["materials"]
    return {"kit_id": kit_id, "material_id": material["id"]}


@pytest.fixture
def scene() -> Scene:
    # Пометка из приказа — не пустая: иначе не видно, что она печатается
    return Scene(marker="**")


# -- ручка: ответ и имя файла ---------------------------------------------


def test_file_is_a_docx_attachment_named_by_topic(client, sms, ai, scene):
    """Учитель складывает планы в одну папку, поэтому имя — по теме урока.
    Тема кириллическая, и в `filename` её не положить: там запасное из ASCII,
    а настоящее имя — в `filename*` UTF-8."""
    ids = ready_material(client, sms, ai, scene)

    resp = download(client, ids["material_id"])

    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == DOCX
    disposition = resp.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    fallback = re.search(r'filename="([^"]*)"', disposition).group(1)
    assert fallback == f"lesson_plan_{ids['kit_id']}.docx"
    assert fallback.isascii()
    encoded = re.search(r"filename\*=UTF-8''(\S+)", disposition).group(1)
    assert encoded.isascii()
    assert unquote(encoded) == f"Краткосрочный план — {TOPIC}.docx"
    # Байты — настоящий документ Word, а не что-то с правильным заголовком
    opened(resp)


def test_file_name_drops_characters_forbidden_in_file_names(client, sms, ai, scene):
    """Слэш, двоеточие, вопрос и кавычки в теме — обычное дело («Дроби: что
    это?»), но в имени файла Windows их не примет, а кавычка ещё и ломает
    сам заголовок. Они вычищаются, остальная тема остаётся."""
    topic = 'Сложение/вычитание: что это? "Разряды" <10> | 20*2\\3'
    ids = ready_material(client, sms, ai, scene, topic=topic)

    resp = download(client, ids["material_id"])

    assert resp.status_code == 200, resp.text
    disposition = resp.headers["content-disposition"]
    encoded = re.search(r"filename\*=UTF-8''(\S+)", disposition).group(1)
    name = unquote(encoded)
    assert name == "Краткосрочный план — Сложениевычитание что это Разряды 10  2023.docx"
    assert not set('/\\:*?"<>|') & set(name)
    # В документе тема — как её написал учитель: чистится только имя файла
    assert head_value(opened(resp), "topic") == [topic]


# -- каждый блок формы доезжает до документа ------------------------------


def _expected(kind: str) -> tuple[str, list[str]]:
    """Где в документе искать блок и какие строки там должны стоять."""
    section = {"title": "Числа и величины", "subsection": "Операции над числами"}
    if kind == "section":
        return "head", [section["title"], section["subsection"]]
    if kind == "learning_objectives":
        return "head", [f"{CODE}** {OBJECTIVE_TEXT}"]
    value = DRAFT_MANY[kind]
    if kind in ("lesson_goals", "assessment_criteria"):
        return "head", bullets(value)
    if kind == "language_goals":
        return "head", [
            f"{RU['vocabulary']}:",
            *bullets(value["vocabulary"]),
            f"{RU['phrases']}:",
            *bullets(value["phrases"]),
        ]
    if kind in ("values", "cross_curricular_links", "prior_knowledge"):
        return "head", [value]
    if kind in STEP_STAGES:
        return "stage", value
    if kind == "reflection":
        return "closing", bullets(value)
    return "closing", [value]


FORM_BLOCKS = ("section", "learning_objectives", *DRAFT)


def test_form_has_fifteen_blocks():
    """Пятнадцать блоков по контракту: два от сервера и тринадцать от модели.
    Без этой сверки параметризация ниже молча проверяла бы меньше."""
    assert len(FORM_BLOCKS) == len(set(FORM_BLOCKS)) == 15


@pytest.mark.parametrize("kind", FORM_BLOCKS)
def test_every_block_lands_in_its_row(client, sms, ai, scene, kind):
    """Блок, пропавший из Word, учитель заметит только у завуча. Каждый
    элемент — отдельный абзац своей строки формы, в порядке ответа модели."""
    ids = ready_material(client, sms, ai, scene)
    document = opened(download(client, ids["material_id"]))

    place, expected = _expected(kind)
    if place == "head":
        assert head_value(document, kind) == expected
    elif place == "closing":
        assert closing_value(document, kind) == expected
    else:
        rows = stage_rows(document, kind)
        assert len(rows) == len(expected)
        for row, step in zip(rows, expected, strict=True):
            assert cell_lines(row.cells[1]) == [step["teacher"]]
            assert cell_lines(row.cells[2]) == [step["students"]]
            # Пустые ресурсы — пустая ячейка, а не «[]» или прочерк
            assert cell_lines(row.cells[3]) == (step["resources"] or [""])


# -- цель обучения ---------------------------------------------------------


def test_objective_is_printed_verbatim_with_code_and_marker(client, sms, ai, scene):
    """Цель обучения — «что нельзя сломать» № 1: завуч сверяет её с приказом
    буква в букву. Код, пометка сразу за ним, пробел и текст справочника —
    целиком, без пересказа, и ровно один раз."""
    ids = ready_material(client, sms, ai, scene)
    document = opened(download(client, ids["material_id"]))

    line = f"{CODE}** {OBJECTIVE_TEXT}"
    assert head_value(document, "learning_objectives") == [line]
    assert [text for text in all_lines(document) if OBJECTIVE_TEXT in text] == [line]


def test_objective_without_marker_has_no_trailing_stars(client, sms, ai):
    """Пустая пометка — ничего между кодом и текстом, а не «None» и не двойной пробел."""
    ids = ready_material(client, sms, ai, Scene())
    document = opened(download(client, ids["material_id"]))

    assert head_value(document, "learning_objectives") == [f"{CODE} {OBJECTIVE_TEXT}"]


# -- таблица хода урока -----------------------------------------------------


def test_course_table_has_a_row_per_step_and_stage_times(client, sms, ai, scene):
    """Ход урока — главная таблица ҚМЖ: шапка из четырёх колонок и по строке
    на каждый шаг. Время этапа модели не выбирать — оно из формы."""
    ids = ready_material(client, sms, ai, scene)
    table = stages_table(opened(download(client, ids["material_id"])))

    steps = sum(len(DRAFT_MANY[stage]) for stage in STEP_STAGES)
    assert len(table.rows) == 1 + steps == 6
    assert [cell.text for cell in table.rows[0].cells] == [
        RU["stage"],
        RU["teacher_actions"],
        RU["students_actions"],
        RU["resources"],
    ]
    stage_texts = [cell_lines(row.cells[0]) for row in table.rows[1:]]
    assert stage_texts == [
        [RU["stage_beginning"], "0–5 мин"],
        *[[RU["stage_middle"], "5–35 мин"]] * 3,
        [RU["stage_end"], "35–40 мин"],
    ]
    # Шаги — в порядке этапов и в порядке ответа модели
    teachers = [row.cells[1].text for row in table.rows[1:]]
    assert teachers == [step["teacher"] for stage in STEP_STAGES for step in DRAFT_MANY[stage]]


def test_stage_cell_is_merged_vertically(client, sms, ai, scene):
    """На бумаге этап — одна ячейка на все его шаги. Без склейки Word печатает
    «Середина урока, 5–35 мин» трижды, и план выглядит сделанным на коленке."""
    ids = ready_material(client, sms, ai, scene)
    table = stages_table(opened(download(client, ids["material_id"])))

    def vmerge(row) -> str | None:
        first_tc = row._tr.tc_lst[0]
        tc_pr = first_tc.tcPr
        node = None if tc_pr is None else tc_pr.find(qn("w:vMerge"))
        if node is None:
            return None
        return node.get(qn("w:val")) or "continue"

    marks = [vmerge(row) for row in table.rows[1:]]
    # Начало и конец — по шагу, склеивать нечего; середина — три шага
    assert marks == [None, "restart", "continue", "continue", None]
    middle = [row.cells[0]._tc for row in table.rows[2:5]]
    assert middle[0] is middle[1] is middle[2]
    # В продолжениях склейки текста нет: этап и время — один раз
    continuation_text = "".join(
        node.text or "" for row in table.rows[3:5] for node in row._tr.tc_lst[0].iter(qn("w:t"))
    )
    assert continuation_text == ""


def test_course_table_header_repeats_on_every_page(client, sms, ai, scene):
    """Ход урока почти всегда переезжает на вторую страницу — без повтора
    шапки там непонятно, где действия педагога, а где учеников."""
    ids = ready_material(client, sms, ai, scene)
    table = stages_table(opened(download(client, ids["material_id"])))

    def repeats(row) -> bool:
        tr_pr = row._tr.trPr
        return tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None

    assert repeats(table.rows[0])
    assert not any(repeats(row) for row in table.rows[1:])


# -- шапка ҚМЖ -------------------------------------------------------------


def test_header_takes_teacher_name_grade_and_topic(client, sms, ai, scene):
    """ФИО — из профиля в момент скачивания, класс и тема — из комплекта.
    Дата урока и присутствующие — строки формы с пустыми ячейками: учитель
    впишет их от руки, а число детей из заказа — не число присутствующих."""
    ids = ordered_by_named(
        client,
        sms,
        ai,
        scene,
        last_name="Нурланова",
        first_name="Айгуль",
        middle_name="Серикқызы",
    )
    document = opened(download(client, ids["material_id"]))

    assert head_value(document, "teacher") == ["Нурланова Айгуль Серикқызы"]
    assert head_value(document, "grade") == ["2"]
    assert head_value(document, "topic") == [TOPIC]
    assert head_value(document, "date") == [""]
    assert head_value(document, "present") == [""]
    assert not any("24" in line for line in all_lines(document))


def test_teacher_name_is_read_at_download_time(client, sms, ai, scene):
    """ФИО в содержимое материала не входит — поправил профиль, и следующий
    файл уже с новым именем, без перегенерации."""
    ids = ordered_by_named(client, sms, ai, scene, last_name="Нурланова", first_name="Айгуль")
    assert head_value(opened(download(client, ids["material_id"])), "teacher") == [
        "Нурланова Айгуль"
    ]

    assert client.patch("/me", json={"last_name": "Сапарова"}).status_code == 200

    assert head_value(opened(download(client, ids["material_id"])), "teacher") == [
        "Сапарова Айгуль"
    ]


def test_empty_profile_gives_an_empty_teacher_cell(client, sms, ai, scene):
    """Вход по SMS заводит человека без ФИО — файл всё равно отдаётся,
    строка педагога пустая, а не «None None»."""
    ids = ready_material(client, sms, ai, scene)

    document = opened(download(client, ids["material_id"]))

    assert head_value(document, "teacher") == [""]
    assert not any("None" in line for line in all_lines(document))


# -- документ из текущего content -----------------------------------------


def test_edited_block_shows_up_in_the_next_download(client, sms, ai, scene):
    """Файл собирается при каждом скачивании (журнал L4, запись 6): одна
    забытая инвалидация сохранённого файла отдала бы завучу старую версию."""
    ids = ready_material(client, sms, ai, scene)
    old = DRAFT_MANY["values"]
    assert head_value(opened(download(client, ids["material_id"])), "values") == [old]

    new = "Дружба и взаимопомощь при работе в парах."
    resp = client.put(
        f"/lesson_ai/materials/{ids['material_id']}/blocks/values", json={"content": new}
    )
    assert resp.status_code == 200, resp.text

    document = opened(download(client, ids["material_id"]))
    assert head_value(document, "values") == [new]
    assert not any(old in line for line in all_lines(document))


# -- шрифт и оформление ----------------------------------------------------


def _normal_rpr(document):
    return document.styles["Normal"].element.rPr


def test_normal_style_is_times_new_roman_in_every_slot(client, sms, ai, scene):
    """Word выбирает слот rFonts по диапазону символа: незаполненный слот —
    и кириллица печатается шрифтом темы. Язык текста — русский, иначе
    проверка орфографии подчеркнёт всё."""
    ids = ready_material(client, sms, ai, scene)
    rpr = _normal_rpr(opened(download(client, ids["material_id"])))

    fonts = rpr.find(qn("w:rFonts"))
    assert {slot: fonts.get(qn(f"w:{slot}")) for slot in ("ascii", "hAnsi", "eastAsia", "cs")} == {
        slot: "Times New Roman" for slot in ("ascii", "hAnsi", "eastAsia", "cs")
    }
    assert [node.get(qn("w:val")) for node in rpr.findall(qn("w:lang"))] == ["ru-RU"]


def test_kazakh_plan_is_kazakh_through_and_through(client, sms, ai):
    """Язык документа — язык обучения: у казахского ҚМЖ казахские подписи
    формы и казахский язык текста для проверки орфографии."""
    ids = ready_material(client, sms, ai, Scene(lang="kz"))
    document = opened(download(client, ids["material_id"]))

    lines = all_lines(document)
    assert LABELS["kz"]["title"] in lines
    assert LABELS["kz"]["course"] in lines
    assert RU["title"] not in lines
    rpr = _normal_rpr(document)
    assert [node.get(qn("w:val")) for node in rpr.findall(qn("w:lang"))] == ["kk-KZ"]


def test_tables_carry_borders_in_schema_order(client, sms, ai, scene):
    """Рамки — в самой таблице: стиль «Table Grid» в Pages и почтовом
    просмотре не виден. И стоят на своём месте в tblPr: на нарушенном
    порядке схемы Word предлагает «восстановить повреждённый файл»."""
    ids = ready_material(client, sms, ai, scene)
    document = opened(download(client, ids["material_id"]))

    assert len(document.tables) == 4
    for table in document.tables:
        tags = [child.tag for child in table._tbl.tblPr]
        assert tags.count(qn("w:tblBorders")) == 1
        borders = tags.index(qn("w:tblBorders"))
        for later in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"):
            if qn(later) in tags:
                assert borders < tags.index(qn(later)), later
        edges = table._tbl.tblPr.find(qn("w:tblBorders"))
        assert [edge.tag for edge in edges] == [
            qn(f"w:{edge}") for edge in ("top", "left", "bottom", "right", "insideH", "insideV")
        ]


# -- отказы ----------------------------------------------------------------


def test_file_requires_login(client, sms, ai, scene):
    ids = ready_material(client, sms, ai, scene)
    client.cookies.clear()

    resp = download(client, ids["material_id"])

    assert resp.status_code == 401, resp.text


def test_someone_elses_material_is_not_found(client, sms, ai, scene):
    """Чужой план — ровно тот же ответ, что несуществующий: по отказу
    нельзя узнать, что такой материал вообще есть."""
    ids = ready_material(client, sms, ai, scene)
    login(client, sms, phone=OTHER_PHONE)

    resp = download(client, ids["material_id"])

    assert resp.status_code == 404, resp.text
    assert resp.json() == NOT_FOUND


def test_missing_material_is_not_found(client, sms, ai, scene):
    ids = ready_material(client, sms, ai, scene)

    resp = download(client, ids["material_id"] + 1000)

    assert resp.status_code == 404, resp.text
    assert resp.json() == NOT_FOUND


def test_queued_material_is_not_ready(client, sms, scene):
    """Кнопку «Скачать» экран показывает по статусу, но запрос мог уйти
    раньше: пустой документ вместо отказа ушёл бы завучу."""
    login(client, sms)
    kit = order(client, scene.objective.id).json()
    [material] = kit["materials"]

    resp = download(client, material["id"])

    assert resp.status_code == 409, resp.text
    assert resp.json() == NOT_READY


def test_failed_material_is_not_ready(client, sms, ai, scene):
    """Три отказа провайдера — материал упал, содержимого у него нет."""
    login(client, sms)
    kit_id = order(client, scene.objective.id).json()["id"]
    ai.fails().fails().fails()
    assert run_next(ai)
    hurry_retries()
    assert run_next(ai)
    hurry_retries()
    assert run_next(ai)
    [material] = client.get(f"/lesson_ai/kits/{kit_id}").json()["materials"]
    assert material["status"] == "failed"

    resp = download(client, material["id"])

    assert resp.status_code == 409, resp.text
    assert resp.json() == NOT_READY


# -- рендер как чистая функция --------------------------------------------


HEADER = PlanHeader(lang="ru", teacher_name="", grade=2, topic="")


def test_render_survives_empty_content():
    """Рендер не должен падать на неполном содержимом: иначе один пропавший
    ключ после ручной правки в базе — 500 вместо документа."""
    document = docx.Document(BytesIO(render_lesson_plan({}, HEADER)))

    assert RU["title"] in all_lines(document)
    # Этапов нет — в ходе урока одна шапка
    assert len(stages_table(document).rows) == 1


def test_render_skips_a_stage_without_steps():
    content = copy.deepcopy(DRAFT_MANY)
    del content["stage_end"]

    document = docx.Document(BytesIO(render_lesson_plan(content, HEADER)))

    table = stages_table(document)
    assert len(table.rows) == 1 + 1 + 3
    assert not stage_rows(document, "stage_end")
    assert RU["stage_end"] not in all_lines(document)
