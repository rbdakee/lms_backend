"""Просмотр материала «Жоспара» на сайте — HTML из того самого файла,
который учитель скачивает.

Разбирается готовый `.docx`/`.pptx`, а не содержимое материала: второй рендер
«по данным» разошёлся бы с Word при первой же правке оформления, и учитель
увидел бы на сайте не то, что распечатает. Из файла берётся только то, чем
пользуются наши рендеры (`app/adapters/docx/*`, `app/adapters/pptx/*`):
абзацы с жирным, курсивом и кеглем, выравнивание, таблицы с объединёнными
ячейками, рамки, разрывы страниц; у презентации — фигуры слайдов (плашки, текст,
картинки) и заметки.

Документ не копирует страницу A4, а переливается под ширину экрана: на 390px
телефона лист в масштабе был бы набран шрифтом в 8px. Таблицы держат пропорции
столбцов, широкие прокручиваются вбок. Слайд, наоборот, держит 16:9 и кегль
в долях ширины — на доске он такой же, только крупнее.

Весь текст экранируется: в нём ответы модели и правки учителя. Скриптов
в странице нет, стили встроены — фронт показывает её в `iframe` без скриптов.
"""

import base64
from html import escape
from io import BytesIO

import docx
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

# Заметки докладчика — как их называет PowerPoint на языке обучения
NOTES = {"ru": "Заметки", "kz": "Ескертпелер"}

# Кегль основного текста документа — 12 pt (`docx/common.py`), он и есть 1em.
# Крупнее 3em не растим: 48 pt цифры кубика на телефоне шире своей клетки
BASE_PT = 12
MAX_EM = 3.0
# Высота строки в Word бывает во всю страницу (диплом игры) — на экране
# хватает её части, остальное — пустое поле для печати
MAX_ROW_EM = 24.0

BORDER = {"single": "solid", "dashed": "dashed", "dotted": "dotted", "double": "double"}
JC = {"center": "center", "right": "right", "end": "right", "both": "justify"}

# Автонумерация списков: тип из `buAutoNum` шаблона — то, что в PowerPoint
# увидит учитель (у вопроса с вариантами — буквы A. B. C., у шагов — 1. 2. 3.)
NUMERALS = {
    "arabicPeriod": lambda n: f"{n}.",
    "alphaUcPeriod": lambda n: f"{chr(64 + n)}.",
    "alphaLcPeriod": lambda n: f"{chr(96 + n)}.",
}


def _auto_num(kind: str | None, number: int) -> str:
    if kind in ("alphaUcPeriod", "alphaLcPeriod") and not 1 <= number <= 26:
        kind = "arabicPeriod"  # буквы кончились — дальше нумеруем цифрами
    return NUMERALS.get(kind, NUMERALS["arabicPeriod"])(number)

STYLE = """
*{box-sizing:border-box}
html{-webkit-text-size-adjust:100%}
body{margin:0;padding:12px;background:#eef0f3;color:#1a1a1a;
  font:15px/1.4 "Times New Roman",Times,serif;overflow-wrap:break-word;hyphens:auto}
.doc{background:#fff;padding:16px 14px;border-radius:6px;box-shadow:0 1px 3px rgba(0,0,0,.12);
  max-width:760px;margin:0 auto}
p{margin:0;min-height:1.4em}
.page{border:0;border-top:1px dashed #9aa0a8;margin:20px -14px}
.scroll{overflow-x:auto;margin:0 -2px;padding:0 2px}
table{border-collapse:collapse;width:100%;table-layout:fixed}
table.wide{min-width:560px}
td{padding:3px 5px;vertical-align:top}
table.lined>tbody>tr>td{border:1px solid #1a1a1a}
.deck{max-width:760px;margin:0 auto;font-family:Arial,Helvetica,sans-serif}
.num{font-size:12px;color:#6b6b6b;margin:0 0 4px}
.slide{position:relative;container-type:inline-size;aspect-ratio:16/9;background:#fff;
  border-radius:6px;box-shadow:0 1px 3px rgba(0,0,0,.18);overflow:hidden;
  line-height:1.15}
.bar{position:absolute}
.bar p{min-height:0;margin:0}
.mk{white-space:pre}
.notes{background:#fff;border-radius:6px;margin:6px 0 0;padding:8px 10px;font-size:13px;
  color:#3a3a3a;line-height:1.35}
.notes b{display:block;font-size:12px;color:#6b6b6b;margin-bottom:2px}
.gap{height:20px}
@media (min-width:640px){body{padding:20px;font-size:16px}.doc{padding:32px 36px}
  .page{margin:28px -36px}}
"""


def render(fmt: str, data: bytes, lang: str) -> str:
    body = _docx(data) if fmt == "docx" else _pptx(data, lang)
    return (
        f'<!doctype html><html lang="{"kk" if lang == "kz" else "ru"}"><head>'
        '<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
        f"<style>{STYLE}</style></head><body>{body}</body></html>"
    )


# -- Word ----------------------------------------------------------------------


def _docx(data: bytes) -> str:
    body = docx.Document(BytesIO(data)).element.body
    return f'<div class="doc">{_blocks(body)}</div>'


def _blocks(parent) -> str:
    out = []
    for child in parent:
        if child.tag == qn("w:p"):
            out.append(_paragraph(child))
        elif child.tag == qn("w:tbl"):
            out.append(_table(child))
    return "".join(out)


def _on(rpr, tag: str) -> bool:
    """`<w:b/>` включает, `<w:b w:val="0"/>` выключает — рендер пишет оба."""
    el = rpr.find(qn(tag)) if rpr is not None else None
    return el is not None and el.get(qn("w:val")) not in ("0", "false")


def _paragraph(p) -> str:
    ppr = p.find(qn("w:pPr"))
    style = []
    before = False
    if ppr is not None:
        jc = ppr.find(qn("w:jc"))
        if jc is not None and jc.get(qn("w:val")) in JC:
            style.append(f"text-align:{JC[jc.get(qn('w:val'))]}")
        spacing = ppr.find(qn("w:spacing"))
        if spacing is not None:
            for attr, prop in (("w:before", "margin-top"), ("w:after", "margin-bottom")):
                twips = spacing.get(qn(attr))
                if twips and twips.isdigit() and int(twips):
                    style.append(f"{prop}:{_em(int(twips) / 20)}")
        before = _on(ppr, "w:pageBreakBefore")

    # Разрыв страницы внутри абзаца режет его на два: так его и видно в Word
    parts: list[list[str]] = [[]]
    for run in p.iter(qn("w:r")):
        rpr = run.find(qn("w:rPr"))
        css = []
        if _on(rpr, "w:b"):
            css.append("font-weight:700")
        if _on(rpr, "w:i"):
            css.append("font-style:italic")
        sz = rpr.find(qn("w:sz")) if rpr is not None else None
        if sz is not None and (sz.get(qn("w:val")) or "").isdigit():
            css.append(f"font-size:{_em(int(sz.get(qn('w:val'))) / 2)}")
        text = []
        for el in run:
            if el.tag == qn("w:t"):
                text.append(escape(el.text or ""))
            elif el.tag == qn("w:tab"):
                text.append(" ")
            elif el.tag == qn("w:br"):
                if el.get(qn("w:type")) == "page":
                    parts[-1].append(_span("".join(text), css))
                    text = []
                    parts.append([])
                else:
                    text.append("<br>")
        parts[-1].append(_span("".join(text), css))

    attr = f' style="{";".join(style)}"' if style else ""
    html = '<hr class="page">'.join(f"<p{attr}>{''.join(part)}</p>" for part in parts)
    return f'<hr class="page">{html}' if before else html


def _span(text: str, css: list[str]) -> str:
    if not text:
        return ""
    return f'<span style="{";".join(css)}">{text}</span>' if css else text


def _em(pt: float) -> str:
    return f"{min(pt / BASE_PT, MAX_EM):.2f}em"


def _table(tbl) -> str:
    # Сетка — только своя: `iter` зашёл бы и во вложенные таблицы
    grid = [int(c.get(qn("w:w")) or 0) for c in tbl.findall(f"{qn('w:tblGrid')}/{qn('w:gridCol')}")]
    total = sum(grid) or 1
    tbl_pr = tbl.find(qn("w:tblPr"))
    lined = tbl_pr is not None and tbl_pr.find(qn("w:tblBorders")) is not None

    # Ячейка, объединённая по вертикали, в Word — стопка ячеек с `vMerge`;
    # в HTML это одна ячейка с `rowspan`, продолжения выбрасываются
    rows = []
    for tr in tbl.findall(qn("w:tr")):
        cells, col = [], 0
        for tc in tr.findall(qn("w:tc")):
            tc_pr = tc.find(qn("w:tcPr"))
            span_el = tc_pr.find(qn("w:gridSpan")) if tc_pr is not None else None
            span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
            merge_el = tc_pr.find(qn("w:vMerge")) if tc_pr is not None else None
            merge = None if merge_el is None else (merge_el.get(qn("w:val")) or "continue")
            cells.append({"tc": tc, "col": col, "span": span, "merge": merge, "rowspan": 1})
            col += span
        rows.append((tr, cells))
    for number, (_, cells) in enumerate(rows):
        for cell in cells:
            if cell["merge"] != "restart":
                continue
            for _, below in rows[number + 1 :]:
                same = next((c for c in below if c["col"] == cell["col"]), None)
                if same is None or same["merge"] != "continue":
                    break
                cell["rowspan"] += 1

    cols = "".join(f'<col style="width:{w / total * 100:.1f}%">' for w in grid)
    body = []
    for tr, cells in rows:
        tr_css = ""
        tr_pr = tr.find(qn("w:trPr"))
        height = tr_pr.find(qn("w:trHeight")) if tr_pr is not None else None
        if height is not None and (height.get(qn("w:val")) or "").isdigit():
            em = min(int(height.get(qn("w:val"))) / 20 / BASE_PT, MAX_ROW_EM)
            tr_css = f' style="height:{em:.2f}em"'
        tds = "".join(_cell(c) for c in cells if c["merge"] != "continue")
        body.append(f"<tr{tr_css}>{tds}</tr>")

    # Четыре столбца абзацев (ход урока) на телефоне сжались бы в столбики
    # по слову — такая таблица прокручивается вбок. Поле игры и жетоны
    # с короткими словами влезают в экран и так
    texts = [c["tc"] for _, cells in rows for c in cells if c["merge"] != "continue"]
    long = sum(len("".join(tc.itertext())) for tc in texts) / max(len(texts), 1) > 40
    wide = len(grid) >= 4 and long
    cls = " ".join(c for c in ("lined" if lined else "", "wide" if wide else "") if c)
    table = (
        f'<table class="{cls}"><colgroup>{cols}</colgroup><tbody>{"".join(body)}</tbody></table>'
    )
    return f'<div class="scroll">{table}</div>'


def _cell(cell: dict) -> str:
    tc_pr = cell["tc"].find(qn("w:tcPr"))
    css = []
    if tc_pr is not None:
        valign = tc_pr.find(qn("w:vAlign"))
        if valign is not None and valign.get(qn("w:val")) in ("center", "bottom"):
            css.append(
                f"vertical-align:{'middle' if valign.get(qn('w:val')) == 'center' else 'bottom'}"
            )
        borders = tc_pr.find(qn("w:tcBorders"))
        for edge in borders if borders is not None else ():
            side = edge.tag.rsplit("}", 1)[-1]
            if side not in ("top", "left", "bottom", "right"):
                continue
            kind = BORDER.get(edge.get(qn("w:val")) or "")
            css.append(f"border-{side}:{'1px ' + kind + ' #1a1a1a' if kind else 'none'}")
    attrs = ""
    if cell["span"] > 1:
        attrs += f' colspan="{cell["span"]}"'
    if cell["rowspan"] > 1:
        attrs += f' rowspan="{cell["rowspan"]}"'
    if css:
        attrs += f' style="{";".join(css)}"'
    return f"<td{attrs}>{_blocks(cell['tc'])}</td>"


# -- PowerPoint ----------------------------------------------------------------


def _pptx(data: bytes, lang: str) -> str:
    """Слайд — рамка 16:9, в ней каждая фигура файла на своём месте (доли
    ширины и высоты слайда), кегль — в долях ширины (`cqw`). Тот же файл,
    что скачает учитель: плашки, кружки, картинки и порядок наложения — как
    в нём. Анимации в просмотре нет: ответ вопроса и число видны сразу."""
    deck = Presentation(BytesIO(data))
    width, height = deck.slide_width, deck.slide_height
    total = len(deck.slides)
    out = []
    for number, slide in enumerate(deck.slides, start=1):
        items = "".join(_shape(shape, width, height) for shape in slide.shapes)
        out.append(
            f'<p class="num">{number} / {total}</p>'
            f'<div class="slide" style="background:{_background(slide)}">{items}</div>'
        )
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if notes.strip():
            lines = "<br>".join(escape(line) for line in notes.splitlines())
            out.append(f'<div class="notes"><b>{NOTES[lang]}</b>{lines}</div>')
        out.append('<div class="gap"></div>')
    return f'<div class="deck">{"".join(out)}</div>'


def _cqw(emu: int, width: int) -> str:
    return f"{emu / width * 100:.2f}cqw"


def _background(slide) -> str:
    try:
        return f"#{slide.background.fill.fore_color.rgb}"
    except (AttributeError, TypeError, ValueError):
        return "#fff"


def _fill(shape) -> str | None:
    try:
        return f"#{shape.fill.fore_color.rgb}"
    except (AttributeError, TypeError, ValueError):
        return None


def _shape(shape, width: int, height: int) -> str:
    css = [
        f"left:{(shape.left or 0) / width * 100:.2f}%",
        f"top:{(shape.top or 0) / height * 100:.2f}%",
        f"width:{(shape.width or 0) / width * 100:.2f}%",
        f"height:{(shape.height or 0) / height * 100:.2f}%",
    ]
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        mime = shape.image.content_type
        src = base64.b64encode(shape.image.blob).decode()
        alt = escape(shape._element.nvPicPr.cNvPr.get("descr") or "", quote=True)
        return (
            f'<img class="bar" alt="{alt}" src="data:{mime};base64,{src}" style="{";".join(css)}">'
        )
    fill = _fill(shape) if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE else None
    if fill:
        css.append(f"background:{fill}")
        geometry = shape._element.spPr.find(qn("a:prstGeom"))
        kind = geometry.get("prst") if geometry is not None else ""
        if kind == "ellipse":
            css.append("border-radius:50%")
        elif kind == "roundRect":
            # Скругление — доля короткой стороны, как у PowerPoint
            radius = min(shape.adjustments[0], 0.5) * min(shape.width, shape.height)
            css.append(f"border-radius:{_cqw(int(radius), width)}")
    if not (shape.has_text_frame and shape.text_frame.text.strip()):
        return f'<div class="bar" style="{";".join(css)}"></div>'
    frame = shape.text_frame
    anchor = {MSO_ANCHOR.MIDDLE: "center", MSO_ANCHOR.BOTTOM: "flex-end"}.get(
        frame.vertical_anchor, "flex-start"
    )
    pad = [frame.margin_top, frame.margin_right, frame.margin_bottom, frame.margin_left]
    css += [
        "display:flex",
        "flex-direction:column",
        f"justify-content:{anchor}",
        "padding:" + " ".join(_cqw(p or 0, width) for p in pad),
    ]
    return f'<div class="bar" style="{";".join(css)}">{_frame(frame, width)}</div>'


def _frame(frame, width: int) -> str:
    align = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right", PP_ALIGN.JUSTIFY: "justify"}
    out = []
    count = 0
    for paragraph in frame.paragraphs:
        css = []
        if paragraph.alignment in align:
            css.append(f"text-align:{align[paragraph.alignment]}")
        if paragraph.space_after is not None:
            css.append(f"margin-bottom:{_cqw(paragraph.space_after, width)}")
        ppr = paragraph._p.pPr
        marker = ""
        if ppr is not None:
            auto = ppr.find(qn("a:buAutoNum"))
            bullet = ppr.find(qn("a:buChar"))
            if auto is not None:
                count += 1
                marker = f"{_auto_num(auto.get('type'), count)} "
            elif bullet is not None:
                # Маркер — содержимое файла, в HTML идёт экранированным
                marker = f"{escape(bullet.get('char') or '\u2022')} "
        runs = []
        for run in paragraph.runs:
            font = run.font
            run_css = []
            if font.size is not None:
                # Подпись этапа в 14 pt на телефоне вышла бы в 5px — ниже 11px не ужимаем
                run_css.append(f"font-size:max(11px,{_cqw(font.size, width)})")
            if font.bold:
                run_css.append("font-weight:700")
            try:
                run_css.append(f"color:#{font.color.rgb}")
            except AttributeError:
                pass
            runs.append(_span(escape(run.text), run_css))
        if marker and runs:
            # Маркер — того же кегля, что первый кусок текста абзаца
            size = paragraph.runs[0].font.size
            style = f' style="font-size:{_cqw(size, width)}"' if size is not None else ""
            runs.insert(0, f'<span class="mk"{style}>{marker}</span>')
        attr = f' style="{";".join(css)}"' if css else ""
        out.append(f"<p{attr}>{''.join(runs)}</p>")
    return "".join(out)
