"""Сборка библиотеки шаблонов слайдов «Жоспара» (PPTX_TEMPLATES_SPEC.md, раздел 1).

Справка, а не часть сервиса: рантайм читает только готовые `library.pptx`
и `catalog.toml`. Геометрия каждого шаблона описана здесь один раз, и из неё
же рисуются `.pptx` (python-pptx) и HTML для снимков `previews/` — так снимок
не расходится с файлом. LibreOffice для снимков не нужен.

    uv run python build_library.py measure OUT   # HTML для замера вместимости
    node shoot.cjs measure OUT                   # -> OUT/measure.json
    uv run python build_library.py build OUT     # library.pptx, catalog.toml, HTML снимков
    node shoot.cjs shots OUT                     # -> OUT/previews/*.png
"""

import html
import json
import math
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

W, H = 13.333, 7.5
IMAGES = Path(sys.argv[3]).resolve() if len(sys.argv) > 3 else None

# Палитра: зелёный луг — основной, солнечно-жёлтый — акцент, остальное
# нейтральное. Красный — только «неверно», смысловой, а не украшение.
G = "1F7A5C"  # основной
G2 = "2A8C6B"  # основной светлее — дальняя волна
GL = "E6F3EC"  # мятный фон и плашки
GM = "9CCBB5"  # рамка места под картинку
Y = "FFC233"  # акцент
YL = "FFF0C2"  # светлый акцент
INK = "1C2B33"  # текст
MUTED = "5B6B73"
WHITE = "FFFFFF"

# Заголовки — Verdana (широкая, открытые формы, различает I/l/1 — для доски
# с пяти метров), текст — Arial. Обе есть на Windows и macOS и рисуют все
# казахские буквы; запасная для Verdana — Tahoma, затем Arial.
FONTS = {
    "V": ("Verdana", "020B0604030504040204", "Verdana, Tahoma, Arial, sans-serif"),
    "A": ("Arial", "020B0604020202020204", "Arial, Helvetica, sans-serif"),
}

LABEL_PIC = "картинка"


# ---------- описание фигур ----------


def tx(
    paras,
    font="A",
    size=28,
    bold=False,
    color=INK,
    align="l",
    anchor="t",
    ins=(7, 4, 7, 4),
    bullet=None,
    mar_l=0,
    indent=0,
    spc_bef=0,
    spc_aft=0,
):
    return dict(
        paras=paras if isinstance(paras, list) else [paras],
        font=font,
        size=size,
        bold=bold,
        color=color,
        align=align,
        anchor=anchor,
        ins=ins,
        bullet=bullet,
        mar_l=mar_l,
        indent=indent,
        spc_bef=spc_bef,
        spc_aft=spc_aft,
    )


def shape(name, geom, x, y, w, h, fill=None, line=None, r=0.0, text=None):
    return dict(
        t="shape", name=name, geom=geom, x=x, y=y, w=w, h=h, fill=fill, line=line, r=r, text=text
    )


def wave(name, y0, amp, periods, phase, fill, n=96):
    pts = [
        (W * i / n, y0 + amp * math.sin(2 * math.pi * periods * i / n + phase))
        for i in range(n + 1)
    ]
    pts += [(W, H), (0, H)]
    return dict(t="free", name=name, points=pts, fill=fill)


def pic(name, x, y, w, h, key):
    return dict(t="pic", name=name, x=x, y=y, w=w, h=h, key=key)


def group(name, children):
    return dict(t="group", name=name, children=children)


def image_slot(name, x, y, s, required=False):
    """Квадратное место под картинку каталога (192x192, прозрачный фон)."""
    el = shape(
        f"slot:{name}",
        "rect",
        x,
        y,
        s,
        s,
        fill=None,
        line=(GM, 1.5, "dash"),
        text=tx(LABEL_PIC, size=16, color=MUTED, align="c", anchor="m"),
    )
    el["slot"] = dict(name=name, kind="image", required=required)
    return el


def text_slot(
    name,
    x,
    y,
    w,
    h,
    text,
    kind="text",
    required=True,
    max_items=None,
    geom="rect",
    fill=None,
    r=0.0,
    line=None,
):
    el = shape(f"slot:{name}", geom, x, y, w, h, fill=fill, r=r, line=line, text=text)
    el["slot"] = dict(name=name, kind=kind, required=required, max_items=max_items)
    return el


def bullets(items, size=28, color=INK, mark=G, auto=None, gap=14, **kw):
    b = {"auto": auto, "color": mark} if auto else {"char": "●", "color": mark}
    mar_l = size * 1.6 if auto else size * 1.25
    return tx(
        items, size=size, color=color, bullet=b, mar_l=mar_l, indent=-mar_l, spc_bef=gap, **kw
    )


def title_slot(text, color=INK, w=12.13, x=0.6, required=True):
    return text_slot(
        "title",
        x,
        0.42,
        w,
        1.05,
        tx(text, font="V", size=36, bold=True, color=color, anchor="m"),
        required=required,
    )


# ---------- шаблоны ----------


def template(key, title, purpose, stage, bg, els, anim=(), notes=""):
    return dict(
        key=key,
        title=title,
        purpose=purpose,
        stage=stage,
        bg=bg,
        els=els,
        anim=list(anim),
        notes=notes,
    )


TEMPLATES = [
    template(
        "title",
        "Титульный",
        "Первый слайд: тема урока, предмет и класс, картинка темы",
        ["beginning"],
        G,
        [
            wave("deco:wave_back", 5.95, 0.22, 1.25, 0.6, G2),
            wave("deco:wave_front", 6.45, 0.18, 1.6, 2.2, GL),
            shape("deco:sun", "oval", 11.55, 0.5, 1.25, 1.25, fill=Y),
            shape("deco:picture_bg", "oval", 8.55, 1.1, 4.1, 4.1, fill=WHITE),
            image_slot("picture", 9.35, 1.9, 2.5),
            text_slot(
                "title",
                0.8,
                1.0,
                7.4,
                3.2,
                tx("Тема урока", font="V", size=48, bold=True, color=WHITE, anchor="b"),
            ),
            text_slot(
                "subtitle",
                0.8,
                4.35,
                7.4,
                1.25,
                tx("Предмет · 3 класс", size=28, color=GL),
                required=False,
            ),
        ],
    ),
    template(
        "today",
        "Сегодня на уроке",
        "План урока для детей: 3–5 коротких пунктов, что будем делать",
        ["beginning"],
        WHITE,
        [
            shape("deco:panel", "round", 8.75, 0.45, 3.98, 6.6, fill=GL, r=0.45),
            shape("deco:sun", "oval", 11.75, 0.2, 1.2, 1.2, fill=Y),
            image_slot("picture", 9.34, 2.35, 2.8),
            title_slot("Сегодня на уроке", w=7.8),
            text_slot(
                "items",
                0.6,
                1.75,
                7.8,
                5.2,
                bullets(
                    ["Вспомним слова", "Узнаем новое правило", "Поиграем в паре", "Подведём итог"],
                    size=28,
                    auto="arabicPeriod",
                    gap=18,
                ),
                kind="bullets",
                max_items=5,
            ),
        ],
    ),
    template(
        "explain",
        "Объяснение с картинкой",
        "Новое правило или понятие: короткое объяснение, картинка и главное в жёлтой плашке",
        ["middle"],
        WHITE,
        [
            title_slot("Новое правило"),
            text_slot(
                "body",
                0.6,
                1.7,
                7.4,
                3.7,
                tx(
                    "Короткое объяснение правила простыми словами: два-три предложения, "
                    "которые учитель прочитает вслух.",
                    size=28,
                ),
            ),
            text_slot(
                "highlight",
                0.6,
                5.7,
                7.4,
                1.2,
                tx("Главное правило", size=28, bold=True, anchor="m", ins=(18, 6, 18, 6)),
                required=False,
                geom="round",
                fill=Y,
                r=0.3,
            ),
            shape("deco:picture_bg", "round", 8.45, 1.7, 4.28, 5.2, fill=GL, r=0.4),
            image_slot("picture", 8.99, 2.1, 3.2),
            text_slot(
                "caption",
                8.6,
                5.55,
                3.98,
                1.2,
                tx("Подпись к картинке", size=24, align="c", anchor="m"),
                required=False,
            ),
        ],
        anim=["highlight"],
    ),
    template(
        "compare",
        "Два столбца",
        "Сравнение двух групп или правил: заголовок и 2–4 пункта в каждом столбце",
        ["middle"],
        WHITE,
        [
            title_slot("Сравним"),
            shape("", "round", 0.6, 1.7, 5.9, 5.2, fill=GL, r=0.35),
            shape("", "round", 6.83, 1.7, 5.9, 5.2, fill=YL, r=0.35),
            image_slot("left_picture", 0.95, 2.0, 1.1),
            text_slot(
                "left_title",
                2.2,
                2.0,
                4.1,
                1.1,
                tx("Первое", font="V", size=28, bold=True, color=G, anchor="m"),
            ),
            text_slot(
                "left_items",
                0.95,
                3.3,
                5.3,
                3.4,
                bullets(["Первый пункт", "Второй пункт", "Третий пункт"], size=26),
                kind="bullets",
                max_items=4,
            ),
            image_slot("right_picture", 7.18, 2.0, 1.1),
            text_slot(
                "right_title",
                8.43,
                2.0,
                4.1,
                1.1,
                tx("Второе", font="V", size=28, bold=True, color="8A5A00", anchor="m"),
            ),
            text_slot(
                "right_items",
                7.18,
                3.3,
                5.3,
                3.4,
                bullets(["Первый пункт", "Второй пункт", "Третий пункт"], size=26, mark="C98A00"),
                kind="bullets",
                max_items=4,
            ),
        ],
        anim=["right_items"],
    ),
    template(
        "big_fact",
        "Крупный факт",
        "Одно число, слово или формула крупно и короткое пояснение — запомнить",
        ["middle", "end"],
        G,
        [
            shape("deco:sun", "oval", 11.85, 0.3, 1.0, 1.0, fill=Y),
            title_slot("Запомни!", color=WHITE, w=11.0, required=False),
            shape("", "round", 0.6, 1.7, 12.13, 2.9, fill=WHITE, r=0.45),
            text_slot(
                "number",
                0.9,
                1.85,
                11.53,
                2.6,
                tx("100", font="V", size=88, bold=True, color=G, align="c", anchor="m"),
                kind="number",
            ),
            text_slot(
                "caption",
                0.6,
                4.95,
                9.6,
                2.0,
                tx("Короткое пояснение: что это значит и где встречается.", size=30, color=WHITE),
                required=False,
            ),
            shape("deco:picture_bg", "oval", 10.45, 4.7, 2.28, 2.28, fill=WHITE),
            image_slot("picture", 10.79, 5.04, 1.6),
        ],
        anim=["caption"],
    ),
    template(
        "question_choice",
        "Вопрос с вариантами",
        "Проверка понимания: вопрос, 2–4 варианта, ответ раскрывается по клику",
        ["middle", "end"],
        WHITE,
        [
            shape(
                "deco:badge",
                "oval",
                0.6,
                0.62,
                1.0,
                1.0,
                fill=Y,
                text=tx("?", font="V", size=40, bold=True, align="c", anchor="m", ins=(0, 0, 0, 0)),
            ),
            text_slot(
                "question",
                1.85,
                0.4,
                10.88,
                1.45,
                tx("Вопрос к классу?", font="V", size=32, bold=True, anchor="m"),
            ),
            text_slot(
                "options",
                0.6,
                2.3,
                7.6,
                3.6,
                bullets(
                    ["Первый вариант", "Второй вариант", "Третий вариант"],
                    size=30,
                    auto="alphaUcPeriod",
                    gap=16,
                ),
                kind="bullets",
                max_items=4,
            ),
            text_slot(
                "answer",
                0.6,
                6.12,
                7.6,
                0.9,
                tx("Ответ", size=26, bold=True, color=WHITE, anchor="m", ins=(22, 4, 22, 4)),
                required=False,
                geom="round",
                fill=G,
                r=0.45,
            ),
            shape("deco:picture_bg", "round", 8.6, 2.3, 4.13, 4.72, fill=GL, r=0.4),
            image_slot("picture", 9.07, 3.07, 3.2),
        ],
        anim=["options", "answer"],
    ),
    template(
        "true_false",
        "Верно или неверно",
        "Утверждение, класс показывает «верно/неверно», объяснение ответа по клику",
        ["beginning", "middle", "end"],
        GL,
        [
            shape("deco:picture_bg", "round", 0.6, 0.6, 4.4, 4.4, fill=WHITE, r=0.4),
            image_slot("picture", 1.0, 1.0, 3.6),
            text_slot(
                "statement",
                5.35,
                0.6,
                7.38,
                4.4,
                tx(
                    "Утверждение, с которым можно согласиться или нет.",
                    size=36,
                    bold=True,
                    anchor="m",
                    ins=(26, 10, 26, 10),
                ),
                geom="round",
                fill=WHITE,
                r=0.4,
            ),
            pic("deco:true_mark", 0.75, 5.45, 1.45, 1.45, "check"),
            pic("deco:false_mark", 2.55, 5.45, 1.45, 1.45, "cross"),
            text_slot(
                "answer",
                4.4,
                5.5,
                8.33,
                1.35,
                tx(
                    "Объяснение ответа",
                    size=28,
                    bold=True,
                    color=WHITE,
                    anchor="m",
                    ins=(24, 6, 24, 6),
                ),
                required=False,
                geom="round",
                fill=G,
                r=0.4,
            ),
        ],
        anim=["answer"],
    ),
    template(
        "steps",
        "Шаги",
        "Порядок действий: 3–5 пронумерованных шагов алгоритма или инструкции",
        ["middle"],
        WHITE,
        [
            shape("", "rect", 0, 0, 4.6, H, fill=G),
            shape("deco:sun", "oval", 4.0, 6.45, 1.2, 1.2, fill=Y),
            text_slot(
                "title",
                0.6,
                0.6,
                3.4,
                2.6,
                tx("Как сделать", font="V", size=36, bold=True, color=WHITE),
            ),
            shape("deco:picture_bg", "oval", 0.75, 3.75, 3.1, 3.1, fill=WHITE),
            image_slot("picture", 1.2, 4.2, 2.2),
            text_slot(
                "steps",
                5.25,
                0.6,
                7.48,
                6.3,
                bullets(
                    ["Первый шаг", "Второй шаг", "Третий шаг", "Четвёртый шаг"],
                    size=28,
                    auto="arabicPeriod",
                    gap=22,
                    anchor="m",
                ),
                kind="bullets",
                max_items=5,
            ),
        ],
        anim=["steps"],
    ),
]


def word_card(i, x):
    return [
        shape(f"deco:card{i}", "round", x, 1.65, 3.81, 5.25, fill=WHITE, r=0.35),
        image_slot(f"picture{i}", x + 0.6, 1.95, 2.6, required=i == 1),
        text_slot(
            f"word{i}",
            x + 0.15,
            4.65,
            3.51,
            1.0,
            tx("слово", size=32, bold=True, align="c", anchor="m"),
            required=i == 1,
        ),
        text_slot(
            f"translation{i}",
            x + 0.15,
            5.65,
            3.51,
            1.0,
            tx("перевод", size=26, color=G, align="c", anchor="t"),
            required=False,
        ),
    ]


TEMPLATES.append(
    template(
        "word_cards",
        "Карточки слов",
        "1–3 новых слова: слово, картинка, перевод. Пустую карточку убрать (deco:cardN)",
        ["beginning", "middle"],
        GL,
        [title_slot("Новые слова")] + word_card(1, 0.6) + word_card(2, 4.76) + word_card(3, 8.92),
        anim=["translation1", "translation2", "translation3"],
    )
)

# Таблица: строки держатся равным шагом абзацев — отступ сверху ставит первую
# строку в середину ряда, интервал после абзаца равен шагу ряда минус строка.
ROW = 1.0
TROW_SIZE = 26


def table_col(i, x, w, head, items, required):
    line = TROW_SIZE * 1.2
    top = (ROW * 72 - line) / 2
    return [
        shape(f"deco:col{i}_bg", "round", x, 1.65, w, 5.3, fill=GL, r=0.3),
        text_slot(
            f"col{i}_head",
            x,
            1.65,
            w,
            0.95,
            tx(head, size=26, bold=True, color=WHITE, align="c", anchor="m"),
            required=required,
            geom="round",
            fill=G,
            r=0.3,
        ),
        text_slot(
            f"col{i}_items",
            x,
            2.75,
            w,
            4.1,
            tx(items, size=TROW_SIZE, align="c", ins=(7, top, 7, 0), spc_aft=ROW * 72 - line),
            kind="bullets",
            required=required,
            max_items=4,
        ),
    ]


TEMPLATES.append(
    template(
        "table",
        "Таблица",
        "Таблица в 2–3 столбца и до 4 строк: короткие ячейки в одну строку. "
        "Для двух столбцов убрать col3 и deco:col3_bg",
        ["middle"],
        WHITE,
        [title_slot("Таблица")]
        + table_col(1, 0.6, 3.95, "Столбец 1", ["строка 1", "строка 2", "строка 3"], True)
        + table_col(2, 4.69, 3.95, "Столбец 2", ["строка 1", "строка 2", "строка 3"], True)
        + table_col(3, 8.78, 3.95, "Столбец 3", ["строка 1", "строка 2", "строка 3"], False),
    )
)

# «Найди пару»: ряды-плашки за текстом, по группе на ряд, чтобы лишние ряды
# убирались одним именем.
PITCH, PILL = 1.25, 0.95
PAIR_SIZE = 28


def pair_row(i):
    y = 1.8 + (i - 1) * PITCH
    return group(
        f"deco:row{i}",
        [
            shape("", "round", 0.6, y, 4.9, PILL, fill=GL, r=PILL / 2),
            shape("", "round", 7.83, y, 4.9, PILL, fill=YL, r=PILL / 2),
            shape("", "oval", 5.28, y + PILL / 2 - 0.13, 0.26, 0.26, fill=G),
            shape("", "oval", 7.79, y + PILL / 2 - 0.13, 0.26, 0.26, fill="C98A00"),
        ],
    )


def pair_col(name, x, items):
    line = PAIR_SIZE * 1.2
    return text_slot(
        name,
        x,
        1.8,
        4.9,
        4.95,
        tx(
            items,
            size=PAIR_SIZE,
            ins=(24, (PILL * 72 - line) / 2, 18, 0),
            spc_aft=PITCH * 72 - line,
        ),
        kind="bullets",
        max_items=4,
    )


TEMPLATES.append(
    template(
        "match_pairs",
        "Найди пару",
        "Соединить пары: 3–4 пункта слева и перемешанные ответы справа. Лишние ряды — deco:rowN",
        ["middle", "end"],
        WHITE,
        [title_slot("Найди пару")]
        + [pair_row(i) for i in range(1, 5)]
        + [
            pair_col("left", 0.6, ["один", "два", "три", "четыре"]),
            pair_col("right", 7.83, ["two", "four", "one", "three"]),
        ],
    )
)

TEMPLATES += [
    template(
        "task_answer",
        "Задание с местом для ответа",
        "Задание классу; рамка для ответа на доске, ответ появляется по клику",
        ["middle"],
        WHITE,
        [
            title_slot("Задание", w=8.2),
            text_slot(
                "task",
                0.6,
                1.55,
                8.2,
                2.6,
                tx("Условие задания: что нужно сделать и с чем.", size=30),
            ),
            shape("deco:picture_bg", "oval", 9.25, 0.55, 3.48, 3.48, fill=GL),
            image_slot("picture", 9.74, 1.04, 2.5),
            shape(
                "deco:answer_box",
                "round",
                0.6,
                4.4,
                12.13,
                2.6,
                fill=WHITE,
                line=(G, 2.5, "dash"),
                r=0.35,
            ),
            text_slot(
                "answer",
                1.0,
                4.6,
                11.33,
                2.2,
                tx("Ответ", size=32, bold=True, color=G, anchor="m"),
                required=False,
            ),
        ],
        anim=["answer"],
    ),
    template(
        "physical",
        "Физминутка",
        "Короткая разминка: 2–4 движения, которые дети повторяют за учителем",
        ["middle"],
        Y,
        [
            wave("deco:wave", 6.55, 0.18, 1.4, 1.0, YL),
            shape("deco:picture_bg", "oval", 8.15, 0.9, 4.4, 4.4, fill=WHITE),
            shape("deco:dot", "oval", 7.55, 4.75, 0.9, 0.9, fill=G),
            image_slot("picture", 8.85, 1.6, 3.0),
            text_slot(
                "title",
                0.6,
                0.55,
                7.2,
                1.4,
                tx("Физминутка", font="V", size=48, bold=True, anchor="m"),
            ),
            text_slot(
                "moves",
                0.6,
                2.15,
                7.0,
                4.2,
                bullets(
                    ["Первое движение", "Второе движение", "Третье движение"],
                    size=30,
                    mark=G,
                    gap=14,
                ),
                kind="bullets",
                max_items=4,
            ),
        ],
    ),
    template(
        "reflection",
        "Рефлексия",
        "Итог урока: 2–3 начала фраз для детей и шкала настроения",
        ["end"],
        G,
        [
            title_slot("Подведём итог", color=WHITE),
            text_slot(
                "prompts",
                0.6,
                1.7,
                8.0,
                5.2,
                bullets(
                    ["Сегодня я узнал…", "Мне было трудно…", "Я могу…"],
                    size=30,
                    mark=G,
                    gap=20,
                    anchor="m",
                    ins=(30, 12, 24, 12),
                ),
                kind="bullets",
                max_items=3,
                geom="round",
                fill=WHITE,
                r=0.4,
            ),
            shape("deco:mood_bg", "round", 8.95, 1.7, 3.78, 5.2, fill=Y, r=0.4),
            image_slot("mood1", 10.17, 2.0, 1.35),
            image_slot("mood2", 10.17, 3.63, 1.35),
            image_slot("mood3", 10.17, 5.26, 1.35),
        ],
    ),
    template(
        "homework",
        "Домашнее задание",
        "Домашнее задание: 1–3 пункта, страница или номер упражнения",
        ["end"],
        GL,
        [
            title_slot("Домашнее задание", w=8.6),
            text_slot(
                "tasks",
                0.6,
                1.7,
                8.4,
                5.2,
                bullets(
                    ["Первое задание", "Второе задание"], size=30, gap=18, ins=(30, 18, 24, 12)
                ),
                kind="bullets",
                max_items=3,
                geom="round",
                fill=WHITE,
                r=0.4,
            ),
            text_slot(
                "page",
                9.4,
                1.7,
                3.33,
                1.55,
                tx("с. 24", font="V", size=40, bold=True, align="c", anchor="m"),
                kind="number",
                required=False,
                geom="round",
                fill=Y,
                r=0.4,
            ),
            shape("deco:picture_bg", "oval", 9.415, 3.6, 3.3, 3.3, fill=WHITE),
            image_slot("picture", 9.865, 4.05, 2.4),
        ],
    ),
]

# ---------- заполненные примеры (только для снимков) ----------

FILLED = {
    "title": {
        "title": "Verb to be",
        "subtitle": "Английский язык · 3 класс",
        "picture": "speaking_head",
    },
    "today": {
        "items": [
            "Вспомним слова о себе",
            "Узнаем формы am, is, are",
            "Поиграем «Кто это?»",
            "Расскажем о друге",
        ],
        "picture": "target",
    },
    "explain": {
        "title": "Глагол to be",
        "body": "To be значит «быть». В настоящем времени у него "
        "три формы: am, is и are. Форма зависит от того, о ком мы говорим.",
        "highlight": "I am · He is · We are",
        "picture": "light_bulb",
        "caption": "Три формы to be",
    },
    "compare": {
        "title": "am, is или are?",
        "left_title": "Один",
        "left_items": ["I am a pupil.", "She is happy.", "It is a cat."],
        "left_picture": "girl",
        "right_title": "Много",
        "right_items": ["We are friends.", "You are tall.", "They are at school."],
        "right_picture": "hugging",
    },
    "big_fact": {
        "title": "Запомни!",
        "number": "I am",
        "caption": "С местоимением I всегда am: I am eight. I am a pupil.",
        "picture": "star",
    },
    "question_choice": {
        "question": "My cat ___ black.",
        "options": ["am", "is", "are"],
        "answer": "is — My cat is black.",
        "picture": "cat",
    },
    "true_false": {
        "statement": "We is friends.",
        "answer": "Неверно: We are friends.",
        "picture": "hugging",
    },
    "steps": {
        "title": "Как задать вопрос",
        "steps": [
            "Возьми предложение: He is tall.",
            "Поставь is в начало.",
            "Добавь знак вопроса.",
            "Is he tall?",
        ],
        "picture": "puzzle",
    },
    "word_cards": {
        "title": "New words",
        "word1": "pupil",
        "translation1": "ученик",
        "picture1": "student",
        "word2": "teacher",
        "translation2": "учитель",
        "picture2": "teacher",
        "word3": "friend",
        "translation3": "друг",
        "picture3": "hugging",
    },
    "table": {
        "title": "Кто? + to be",
        "col1_head": "Кто?",
        "col2_head": "to be",
        "col3_head": "Пример",
        "col1_items": ["I", "He / She / It", "We / You", "They"],
        "col2_items": ["am", "is", "are", "are"],
        "col3_items": ["I am happy.", "She is ten.", "We are pupils.", "They are cats."],
    },
    "match_pairs": {
        "left": ["I", "She", "We", "They"],
        "right": ["are pupils.", "am eight.", "is at home.", "are friends."],
    },
    "task_answer": {
        "task": "Вставь am, is или are: My name ___ Aru. I ___ eight. We ___ in class 3.",
        "answer": "is · am · are",
        "picture": "pencil",
    },
    "physical": {
        "moves": [
            "I am a frog — прыгаем!",
            "I am a bird — машем!",
            "I am a tree — тянемся!",
            "I am a cat — садимся.",
        ],
        "picture": "frog",
    },
    "reflection": {
        "prompts": ["Я знаю формы am, is, are.", "Мне было трудно…", "Я могу рассказать о друге."],
        "mood1": "happy",
        "mood2": "thinking",
        "mood3": "sad",
    },
    "homework": {
        "tasks": ["Напиши 5 предложений о семье с am, is, are.", "Выучи новые слова."],
        "page": "с. 24",
        "picture": "notebook",
    },
}
FILLED_KZ = (
    "title",
    {
        "title": "Әдебиеттік оқу",
        "subtitle": "Ы. Алтынсарин «Кел, балалар, оқылық!» · Қазақ тілі, 3 сынып",
        "picture": "open_book",
    },
)


# ---------- PPTX ----------


def emu(v):
    return Emu(int(round(v * 914400)))


def strip_style(shp):
    # Стиль темы по умолчанию даёт тень и белый текст — всё задаём явно
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)


def set_font(run_or_rpr, font_key, size, bold, color, lang="ru-RU"):
    name, panose, _ = FONTS[font_key]
    rpr = run_or_rpr
    rpr.set("lang", lang)
    rpr.set("sz", str(int(size * 100)))
    rpr.set("b", "1" if bold else "0")
    for tag in ("a:solidFill", "a:latin", "a:cs"):
        for old in rpr.findall(qn(tag)):
            rpr.remove(old)
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", color)
    lat = etree.SubElement(rpr, qn("a:latin"))
    lat.set("typeface", name)
    lat.set("panose", panose)
    cs = etree.SubElement(rpr, qn("a:cs"))
    cs.set("typeface", name)


def fill_text(shp, t, values=None):
    tf = shp.text_frame
    tf.word_wrap = True
    # Подбор кегля исполнитель делает сам (лимиты каталога); normAutofit — чтобы
    # при ручной правке в PowerPoint текст ужимался, а не вылезал за рамку
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    l_, t_, r_, b_ = t["ins"]
    tf.margin_left, tf.margin_top = Pt(l_), Pt(t_)
    tf.margin_right, tf.margin_bottom = Pt(r_), Pt(b_)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[
        t["anchor"]
    ]
    paras = values if values is not None else t["paras"]
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[t["align"]]
        ppr = p._p.get_or_add_pPr()
        if t["mar_l"]:
            ppr.set("marL", str(int(Pt(t["mar_l"]))))
            ppr.set("indent", str(int(Pt(t["indent"]))))
        if t["spc_bef"]:
            el = etree.SubElement(ppr, qn("a:spc_bef"))
            etree.SubElement(el, qn("a:spcPts")).set("val", str(int(t["spc_bef"] * 100)))
        if t["spc_aft"]:
            el = etree.SubElement(ppr, qn("a:spc_aft"))
            etree.SubElement(el, qn("a:spcPts")).set("val", str(int(t["spc_aft"] * 100)))
        b = t["bullet"]
        if b:
            etree.SubElement(ppr, qn("a:buClr")).append(
                etree.Element(qn("a:srgbClr"), val=b["color"])
            )
            etree.SubElement(ppr, qn("a:buSzPct")).set(
                "val", "100000" if b.get("auto") else "80000"
            )
            fnt = etree.SubElement(ppr, qn("a:buFont"))
            fnt.set("typeface", FONTS[t["font"]][0] if b.get("auto") else "Arial")
            if b.get("auto"):
                etree.SubElement(ppr, qn("a:buAutoNum")).set("type", b["auto"])
            else:
                etree.SubElement(ppr, qn("a:buChar")).set("char", b["char"])
        else:
            etree.SubElement(ppr, qn("a:buNone"))
        r = p.add_run()
        r.text = text
        set_font(r._r.get_or_add_rPr(), t["font"], t["size"], t["bold"], t["color"])
        end = p._p.get_or_add_endParaRPr()
        set_font(end, t["font"], t["size"], t["bold"], t["color"])


def add_el(shapes, el):
    if el["t"] == "group":
        g = shapes.add_group_shape()
        for ch in el["children"]:
            add_el(g.shapes, ch)
        g.name = el["name"]
        return
    if el["t"] == "pic":
        p = shapes.add_picture(
            str(IMAGES / f"{el['key']}.png"), emu(el["x"]), emu(el["y"]), emu(el["w"]), emu(el["h"])
        )
        p.name = el["name"]
        return
    if el["t"] == "free":
        (x0, y0), rest = el["points"][0], el["points"][1:]
        fb = shapes.build_freeform(emu(x0), emu(y0), scale=914400)
        fb.add_line_segments(rest, close=True)
        shp = fb.convert_to_shape(0, 0)
        strip_style(shp)
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(el["fill"])
        shp.line.fill.background()
        shp.name = el["name"]
        return
    geom = {
        "rect": MSO_SHAPE.RECTANGLE,
        "round": MSO_SHAPE.ROUNDED_RECTANGLE,
        "oval": MSO_SHAPE.OVAL,
    }[el["geom"]]
    shp = shapes.add_shape(geom, emu(el["x"]), emu(el["y"]), emu(el["w"]), emu(el["h"]))
    strip_style(shp)
    if el["geom"] == "round":
        shp.adjustments[0] = min(0.5, el["r"] / min(el["w"], el["h"]))
    if el["fill"]:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(el["fill"])
    else:
        shp.fill.background()
    if el["line"]:
        color, width, dash = el["line"]
        shp.line.color.rgb = RGBColor.from_string(color)
        shp.line.width = Pt(width)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    else:
        shp.line.fill.background()
    if el["text"]:
        fill_text(shp, el["text"])
    shp.name = el["name"] or f"panel_{el['geom']}"


def build_pptx(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    for tpl in TEMPLATES:
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(tpl["bg"])
        for el in tpl["els"]:
            add_el(s.shapes, el)
    prs.core_properties.title = "Жоспар — библиотека шаблонов слайдов"
    prs.core_properties.author = "Жоспар"
    prs.save(path)


# ---------- HTML для снимков и замера ----------


def css_text(t):
    _, _, fam = FONTS[t["font"]]
    l_, t_, r_, b_ = t["ins"]
    just = {"t": "flex-start", "m": "center", "b": "flex-end"}[t["anchor"]]
    align = {"l": "left", "c": "center", "r": "right"}[t["align"]]
    return (
        f"display:flex;flex-direction:column;justify-content:{just};"
        f"padding:{t_}pt {r_}pt {b_}pt {l_}pt;font-family:{fam};font-size:{t['size']}pt;"
        f"font-weight:{700 if t['bold'] else 400};color:#{t['color']};text-align:{align};"
        f"line-height:1.2;"
    )


def marker(b, i):
    if not b:
        return ""
    if b.get("auto") == "arabicPeriod":
        return f"{i + 1}."
    if b.get("auto") == "alphaUcPeriod":
        return f"{'ABCDEFGH'[i]}."
    return b["char"]


def html_paras(t, values):
    out = []
    for i, text in enumerate(values):
        b = t["bullet"]
        style = f"position:relative;padding-left:{t['mar_l']}pt;"
        if i or t["spc_bef"]:
            style += f"margin-top:{t['spc_bef']}pt;"
        if t["spc_aft"] and i < len(values) - 1:
            style += f"margin-bottom:{t['spc_aft']}pt;"
        m = ""
        if b:
            size = "100%" if b.get("auto") else "80%"
            m = (
                f"<span style='position:absolute;left:{t['mar_l'] + t['indent']}pt;"
                f"color:#{b['color']};font-size:{size};"
                f"{'' if b.get('auto') else 'line-height:1.5;'}'>{marker(b, i)}</span>"
            )
        out.append(f"<div style='{style}'>{m}{html.escape(text)}</div>")
    return "".join(out)


def box(x, y, w, h):
    return f"position:absolute;left:{x}in;top:{y}in;width:{w}in;height:{h}in;box-sizing:border-box;"


def html_el(el, fill=None, measure=False):
    """fill=None — пустой шаблон; иначе словарь слотов: незаполненный необязательный
    слот исполнитель убирает, так же и здесь."""
    if el["t"] == "group":
        return "".join(html_el(c, fill) for c in el["children"])
    if el["t"] == "pic":
        return (
            f"<img src='{IMAGES / (el['key'] + '.png')}' "
            f"style='{box(el['x'], el['y'], el['w'], el['h'])}'>"
        )
    if el["t"] == "free":
        pts = " ".join(f"{x * 96:.1f},{y * 96:.1f}" for x, y in el["points"])
        return (
            f"<svg style='position:absolute;left:0;top:0' width='1280' height='720'>"
            f"<polygon points='{pts}' fill='#{el['fill']}'/></svg>"
        )
    slot = el.get("slot")
    value = None
    if slot and fill is not None:
        value = fill.get(slot["name"])
        if value is None and not slot["required"]:
            return ""
    style = box(el["x"], el["y"], el["w"], el["h"])
    if el["fill"]:
        style += f"background:#{el['fill']};"
    if el["line"] and not (slot and slot["kind"] == "image" and value):
        color, width, dash = el["line"]
        style += f"border:{width}pt {'dashed' if dash else 'solid'} #{color};"
    if el["geom"] == "oval":
        style += "border-radius:50%;"
    elif el["geom"] == "round":
        style += f"border-radius:{min(el['r'], min(el['w'], el['h']) / 2)}in;"
    if slot and slot["kind"] == "image" and value:
        return (
            f"<div style='{style}'><img src='{IMAGES / (value + '.png')}' "
            f"style='width:100%;height:100%;object-fit:contain'></div>"
        )
    inner = ""
    if el["text"]:
        t = el["text"]
        style += css_text(t)
        values = value if isinstance(value, list) else ([value] if value else t["paras"])
        attr = ""
        if measure and slot and slot["kind"] != "image":
            attr = (
                f" data-slot='{slot['name']}' data-kind='{slot['kind']}' "
                f"data-items='{slot['max_items'] or 1}' "
                f"data-oneline='{1 if el['text']['spc_aft'] or slot['kind'] == 'number' else 0}'"
            )
        inner = f"<div class='in'{attr}>{html_paras(t, values)}</div>"
    return f"<div style='{style}'>{inner}</div>"


def slide_html(tpl, fill=None, measure=False):
    body = "".join(html_el(el, fill, measure) for el in tpl["els"])
    return (
        f"<div class='slide' data-key='{tpl['key']}' style='position:relative;width:1280px;"
        f"height:720px;overflow:hidden;background:#{tpl['bg']}'>{body}</div>"
    )


PAGE = "<!doctype html><meta charset='utf-8'><style>body{{margin:0}}.in{{width:100%}}</style>{}"


# ---------- каталог ----------


def round_limit(n):
    n = int(n * 0.9)  # запас: PowerPoint переносит строки чуть иначе, чем браузер
    return n - n % 5 if n >= 20 else n


def catalog(limits):
    out = [CATALOG_HEAD]
    for i, tpl in enumerate(TEMPLATES, 1):
        # Картинки в конце: модели сначала нужен текст слайда
        slots = sorted(
            (el["slot"] for el in tpl["els"] if el.get("slot")),
            key=lambda s: s["kind"] == "image",
        )
        decos = []
        for el in tpl["els"]:
            if el.get("name", "").startswith("deco:"):
                decos.append(el["name"][5:])
        out.append("[[template]]")
        out.append(f'key = "{tpl["key"]}"')
        out.append(f"slide = {i}")
        out.append(f'title = "{tpl["title"]}"')
        out.append(f'purpose = "{tpl["purpose"]}"')
        out.append("stage = [" + ", ".join(f'"{s}"' for s in tpl["stage"]) + "]")
        for s in slots:
            out.append("  [[template.slot]]")
            out.append(f'  name = "{s["name"]}"')
            out.append(f'  kind = "{s["kind"]}"')
            if s["kind"] != "image":
                out.append(f"  max_chars = {limits[tpl['key']][s['name']]}")
            if s["kind"] == "bullets":
                out.append(f"  max_items = {s['max_items']}")
            out.append(f"  required = {'true' if s['required'] else 'false'}")
        out.append("  [template.deco]")
        out.append("  removable = [" + ", ".join(f'"{d}"' for d in decos) + "]")
        out.append("  [template.animation]")
        out.append("  default = [" + ", ".join(f'"{a}"' for a in tpl["anim"]) + "]")
        out.append("")
    return "\n".join(out)


CATALOG_HEAD = """\
# Каталог шаблонных слайдов «Жоспара» — PPTX_TEMPLATES_SPEC.md, раздел 1.
# Записи идут в порядке слайдов library.pptx; slide — номер слайда с 1.
# Собрано build/build_library.py из одного описания геометрии; правка руками
# в library.pptx без правки скрипта разойдётся со снимками previews/.
#
# Как устроены слоты (для исполнителя):
# - text / number — одна фигура, один абзац с одним run. Формат (шрифт, кегль,
#   цвет, жирность) брать из rPr первого run и вписывать текст в него.
# - bullets — одна фигура, абзац на пункт. Новый пункт — копия pPr/rPr первого
#   абзаца; маркер или автонумерация уже в pPr (buChar / buAutoNum).
#   В table и match_pairs пункты — ряды: шаг ряда держится интервалом после
#   абзаца, поэтому ячейка обязана помещаться в одну строку (max_chars).
# - image — прямоугольник с пунктирной рамкой и подписью «картинка». Его
#   заменяют картинкой каталога, вписанной в рамку с сохранением пропорций
#   (рамки квадратные, картинки каталога 192x192), а сам прямоугольник удаляют;
#   имя slot:<name> лучше перенести на картинку.
# - Незаполненный необязательный слот исполнитель удаляет целиком: в шаблоне
#   стоит текст-заполнитель («Тема урока», «картинка»), в комплект он попасть
#   не должен. У answer/highlight/statement/prompts/tasks/page плашка — это
#   сама фигура слота, она уходит вместе с ним.
# - Кегль: в шаблоне стоит номинальный (текст 26–32 pt, заголовки 36–48 pt,
#   number 40–88 pt). max_chars — вместимость при номинальном кегле с запасом
#   10 % (замерено в браузере теми же шрифтами, для bullets — при max_items
#   пунктах). Превышение — уменьшать кегль, но не ниже 24 pt, и предупреждать.
#   В фигурах включено «сжать текст при переполнении» (normAutofit) — это для
#   ручной правки в PowerPoint; при открытии PowerPoint его не пересчитывает,
#   кегль ставит исполнитель.
# - deco:* — украшения: фигура, картинка (true_mark/false_mark — значки
#   каталога check/cross) или группа (rowN в match_pairs). Удалять можно любую.
#   Если не заполнена картинка, рядом стоящую подложку deco:picture_bg (и
#   deco:cardN в word_cards, deco:col3_bg в table, deco:rowN в match_pairs при
#   меньшем числе пунктов) тоже лучше убрать — сам исполнитель этого не знает,
#   правило стоит в purpose шаблона.
# - Фигуры без префикса — каркас (панели, карточки столбцов), имена panel_*.
#
# Шрифты: заголовки Verdana (запасная Tahoma, затем Arial), текст Arial — есть
# на Windows и macOS, все казахские буквы на месте. Шрифты не встраиваются.
# Палитра: основной 1F7A5C, акцент FFC233, нейтральные E6F3EC / FFF0C2 / 1C2B33.

"""


def main():
    mode, out = sys.argv[1], Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)
    if mode == "measure":
        (out / "measure.html").write_text(
            PAGE.format("".join(slide_html(t, measure=True) for t in TEMPLATES))
        )
        return
    measured = json.loads((out / "measure.json").read_text())
    limits = {k: {s: round_limit(n) for s, n in v.items()} for k, v in measured.items()}
    build_pptx(out / "library.pptx")
    (out / "catalog.toml").write_text(catalog(limits))
    shots = out / "html"
    shots.mkdir(exist_ok=True)
    by_key = {t["key"]: t for t in TEMPLATES}
    for t in TEMPLATES:
        (shots / f"{t['key']}.html").write_text(PAGE.format(slide_html(t)))
        (shots / f"{t['key']}__filled.html").write_text(
            PAGE.format(slide_html(t, fill=FILLED[t["key"]]))
        )
    key, fill = FILLED_KZ
    (shots / f"{key}__filled_kz.html").write_text(PAGE.format(slide_html(by_key[key], fill=fill)))


if __name__ == "__main__":
    main()
